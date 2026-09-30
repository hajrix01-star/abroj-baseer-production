import base64
import csv
from collections import Counter
from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO
from zipfile import BadZipFile, ZipFile

from odoo import _, api, fields, models
from odoo.exceptions import AccessError, UserError, ValidationError


IMPORT_COLUMNS = (
    'category', 'name', 'pricing_method', 'quantity',
    'material_unit_cost', 'auxiliary_unit_cost', 'labor_unit_cost',
    'inclusive_unit_cost', 'lump_sum_cost', 'notes',
    'path', 'node_kind',
)
LEGACY_IMPORT_COLUMNS = IMPORT_COLUMNS[:10]
MAX_IMPORT_ROWS = 500
MAX_IMPORT_SIZE = 5 * 1024 * 1024
MAX_XLSX_UNCOMPRESSED_SIZE = 25 * 1024 * 1024
MAX_XLSX_COMPRESSION_RATIO = 100

PRICING_METHOD_LABELS = {
    'تفصيلي': 'detailed',
    'توريد فقط': 'supply_only',
    'عمل يد فقط': 'labor_only',
    'توريد وتركيب شامل': 'supply_install',
    'مقطوعية': 'lump_sum',
}
VALID_PRICING_METHODS = {
    'detailed', 'supply_only', 'labor_only', 'supply_install', 'lump_sum',
}
NODE_KIND_LABELS = {'قسم': 'section', 'بند': 'item', 'بند مسعّر': 'item'}
VALID_NODE_KINDS = {'section', 'item'}


class AbrojCostPlanImportWizard(models.TransientModel):
    _name = 'abroj.cost.plan.import.wizard'
    _description = 'Import or Export Abroj Project Cost Plan'

    project_id = fields.Many2one(
        'abroj.cost.project', string='المشروع', required=True,
        check_company=True, readonly=True,
    )
    company_id = fields.Many2one(
        'res.company', string='الشركة', required=True,
        default=lambda self: self.env.company, readonly=True, index=True,
    )
    import_file = fields.Binary('ملف دراسة المشروع')
    import_filename = fields.Char('اسم الملف')
    export_file = fields.Binary('ملف التصدير', readonly=True)
    export_filename = fields.Char('اسم ملف التصدير', readonly=True)

    def _ensure_project_scope(self):
        self.ensure_one()
        if self.company_id != self.project_id.company_id:
            raise AccessError(_('يجب أن تتطابق شركة المعالج مع شركة المشروع.'))
        if self.project_id.company_id != self.env.company:
            raise AccessError(_('اختر شركة المشروع نفسها قبل الاستيراد أو التصدير.'))
        if not self.project_id.company_id.abroj_project_costing_enabled:
            raise AccessError(_('تكاليف المشاريع غير مفعلة للشركة المختارة.'))
        if self.project_id.state == 'closed':
            raise UserError(_('لا يمكن استيراد بنود إلى مشروع مغلق.'))

    @api.constrains('project_id', 'company_id')
    def _check_project_company(self):
        for wizard in self:
            if wizard.project_id and wizard.company_id != wizard.project_id.company_id:
                raise ValidationError(_('يجب أن تتطابق شركة المعالج مع شركة المشروع.'))

    @staticmethod
    def _clean_text(value):
        return str(value or '').strip()

    @staticmethod
    def _decimal(value, row_number, column, default=Decimal('0')):
        text = str(value or '').strip().replace(',', '')
        if not text:
            return default
        try:
            amount = Decimal(text)
        except InvalidOperation as error:
            raise ValidationError(_('الصف %(row)s: قيمة غير صالحة في %(column)s.') % {
                'row': row_number, 'column': column,
            }) from error
        if not amount.is_finite() or amount < 0 or amount.as_tuple().exponent < -2:
            raise ValidationError(_('الصف %(row)s: يجب أن تكون %(column)s رقماً غير سالب بمنزلتين عشريتين كحد أقصى.') % {
                'row': row_number, 'column': column,
            })
        return amount

    @classmethod
    def _pricing_method(cls, value, row_number):
        method = cls._clean_text(value)
        method = PRICING_METHOD_LABELS.get(method, method)
        if method not in VALID_PRICING_METHODS:
            raise ValidationError(_('الصف %(row)s: طريقة التسعير غير صالحة. استخدم إحدى القيم الموجودة في النموذج.') % {
                'row': row_number,
            })
        return method

    @staticmethod
    def _header_map(headers):
        normalized = {str(name or '').strip().lower(): index for index, name in enumerate(headers)}
        # path/node_kind were added after the first released template.  Keep
        # accepting that flat template as root priced leaves.
        missing = [column for column in LEGACY_IMPORT_COLUMNS if column not in normalized]
        if missing:
            raise ValidationError(_('الملف لا يحتوي الأعمدة المطلوبة: %(columns)s.') % {
                'columns': ', '.join(missing),
            })
        return normalized

    def _read_csv(self, content):
        try:
            stream = StringIO(content.decode('utf-8-sig'))
            reader = csv.reader(stream)
            headers = next(reader, None)
            if not headers:
                raise ValidationError(_('ملف CSV فارغ.'))
            header_map = self._header_map(headers)
            records = []
            for row_number, row in enumerate(reader, start=2):
                if not any(self._clean_text(cell) for cell in row):
                    continue
                records.append((row_number, {
                    column: row[position] if position < len(row) else ''
                    for column, position in header_map.items()
                }))
                # Stop reading at the limit rather than constructing an
                # arbitrarily large list from a small-but-dense CSV file.
                if len(records) > MAX_IMPORT_ROWS:
                    raise ValidationError(_('الحد الأقصى للاستيراد هو %(count)s بنداً في الملف الواحد.') % {
                        'count': MAX_IMPORT_ROWS,
                    })
            return records
        except UnicodeDecodeError as error:
            raise ValidationError(_('يجب أن يكون ملف CSV محفوظاً بترميز UTF-8.')) from error

    def _read_xlsx(self, content):
        try:
            from openpyxl import load_workbook
        except ImportError as error:
            raise UserError(_('لا تتوفر مكتبة قراءة ملفات Excel في بيئة أودو.')) from error
        try:
            try:
                with ZipFile(BytesIO(content)) as archive:
                    file_infos = archive.infolist()
                    uncompressed_size = sum(info.file_size for info in file_infos)
                    suspicious_member = any(
                        info.file_size > MAX_XLSX_COMPRESSION_RATIO * max(info.compress_size, 1)
                        for info in file_infos
                    )
                    if uncompressed_size > MAX_XLSX_UNCOMPRESSED_SIZE or suspicious_member:
                        raise ValidationError(_(
                            'محتوى ملف Excel بعد فك الضغط أكبر من الحد المسموح.'
                        ))
            except BadZipFile as error:
                raise ValidationError(_('ملف Excel غير صالح.')) from error
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
            sheet = workbook['بنود الدراسة'] if 'بنود الدراسة' in workbook.sheetnames else workbook.active
            if sheet.max_row > MAX_IMPORT_ROWS + 1:
                raise ValidationError(_('الحد الأقصى للاستيراد هو %(count)s بنداً في الملف الواحد.') % {
                    'count': MAX_IMPORT_ROWS,
                })
            rows = sheet.iter_rows(values_only=True)
            headers = next(rows, None)
            if not headers:
                raise ValidationError(_('ملف Excel فارغ.'))
            header_map = self._header_map(headers)
            records = []
            for index, row in enumerate(rows, start=2):
                if not any(self._clean_text(cell) for cell in row):
                    continue
                records.append((index, {
                    column: row[position] if position < len(row) else ''
                    for column, position in header_map.items()
                }))
            return records
        except ValidationError:
            raise
        except Exception as error:
            raise ValidationError(_('تعذر قراءة ملف Excel. استخدم النموذج الذي تم تنزيله من النظام.')) from error

    def _read_import_rows(self):
        self.ensure_one()
        if not self.import_file or not self.import_filename:
            raise ValidationError(_('ارفع ملف دراسة المشروع أولاً.'))
        content = base64.b64decode(self.import_file)
        if not content:
            raise ValidationError(_('ملف دراسة المشروع فارغ.'))
        if len(content) > MAX_IMPORT_SIZE:
            raise ValidationError(_('حجم الملف أكبر من الحد المسموح (5 ميغابايت).'))
        extension = self.import_filename.rsplit('.', 1)[-1].lower() if '.' in self.import_filename else ''
        if extension == 'csv':
            rows = self._read_csv(content)
        elif extension == 'xlsx':
            rows = self._read_xlsx(content)
        else:
            raise ValidationError(_('الصيغ المقبولة هي CSV أو XLSX فقط.'))
        if not rows:
            raise ValidationError(_('لا يحتوي الملف أي بند دراسة.'))
        if len(rows) > MAX_IMPORT_ROWS:
            raise ValidationError(_('الحد الأقصى للاستيراد هو %(count)s بنداً في الملف الواحد.') % {
                'count': MAX_IMPORT_ROWS,
            })
        return rows

    def _validated_values(self):
        self._ensure_project_scope()
        categories = self.env['abroj.cost.category'].search([
            ('company_id', '=', self.project_id.company_id.id), ('active', '=', True),
        ])
        category_by_name = {self._clean_text(category.name): category for category in categories}
        values = []
        identities = []
        known_paths = set()
        for sequence, (row_number, row) in enumerate(self._read_import_rows(), start=10):
            path = self._clean_text(row.get('path'))
            node_kind = NODE_KIND_LABELS.get(self._clean_text(row.get('node_kind')), self._clean_text(row.get('node_kind')) or 'item')
            if node_kind not in VALID_NODE_KINDS:
                raise ValidationError(_('الصف %(row)s: نوع العقدة غير صالح. استخدم section أو item.') % {'row': row_number})
            name = self._clean_text(row['name'])
            if not name:
                raise ValidationError(_('الصف %(row)s: اسم البند مطلوب.') % {'row': row_number})
            path_parts = [part.strip() for part in path.split('/') if part.strip()]
            if path and (not path_parts or path_parts[-1] != name):
                raise ValidationError(_('الصف %(row)s: يجب أن ينتهي المسار باسم العقدة نفسها.') % {'row': row_number})
            if node_kind == 'section' and not path:
                raise ValidationError(_('الصف %(row)s: القسم يحتاج مساراً مثل "المطبخ/سباكة".') % {'row': row_number})
            parent_path = '/'.join(path_parts[:-1]) if path else False
            if parent_path and parent_path not in known_paths:
                raise ValidationError(_('الصف %(row)s: أضف القسم الأب "%(path)s" في صف سابق أولاً.') % {
                    'row': row_number, 'path': parent_path,
                })
            identity = path or 'legacy:%s' % name
            identities.append(identity)
            if path:
                known_paths.add(path)
            row_values = {
                'project_id': self.project_id.id,
                'sequence': sequence,
                'name': name,
                'notes': self._clean_text(row['notes']),
                'node_kind': node_kind,
                '_path': path,
                '_parent_path': parent_path,
            }
            if node_kind == 'item':
                category_name = self._clean_text(row['category'])
                category = category_by_name.get(category_name)
                if not category:
                    raise ValidationError(_('الصف %(row)s: نوع العمل "%(category)s" غير موجود أو غير نشط في الشركة المختارة.') % {
                        'row': row_number, 'category': category_name or '—',
                    })
                quantity = self._decimal(row['quantity'], row_number, 'quantity', default=Decimal('1'))
                if quantity <= 0:
                    raise ValidationError(_('الصف %(row)s: الكمية يجب أن تكون أكبر من صفر.') % {'row': row_number})
                row_values.update({
                    'category_id': category.id,
                    'pricing_method': self._pricing_method(row['pricing_method'], row_number),
                    'quantity': float(quantity),
                    'material_unit_cost': float(self._decimal(row['material_unit_cost'], row_number, 'material_unit_cost')),
                    'auxiliary_unit_cost': float(self._decimal(row['auxiliary_unit_cost'], row_number, 'auxiliary_unit_cost')),
                    'labor_unit_cost': float(self._decimal(row['labor_unit_cost'], row_number, 'labor_unit_cost')),
                    'inclusive_unit_cost': float(self._decimal(row['inclusive_unit_cost'], row_number, 'inclusive_unit_cost')),
                    'lump_sum_cost': float(self._decimal(row['lump_sum_cost'], row_number, 'lump_sum_cost')),
                })
            else:
                row_values.update({'quantity': 0.0, 'pricing_method': 'detailed'})
            values.append(row_values)
        duplicates = [identity for identity, count in Counter(identities).items() if count > 1]
        if duplicates:
            raise ValidationError(_('يتضمن الملف مسارات أو أسماء بنود مكررة: %(names)s.') % {
                'names': ', '.join(duplicates),
            })
        return values

    def _lock_empty_project_plan(self):
        """Serialize imports for one project and recheck its current plan."""
        self.ensure_one()
        self._ensure_project_scope()
        self.env.cr.execute(
            'SELECT id FROM abroj_cost_project WHERE id = %s FOR UPDATE',
            [self.project_id.id],
        )
        if self.env['abroj.cost.plan.line'].search_count([
            ('project_id', '=', self.project_id.id),
        ]):
            raise UserError(_(
                'لا يمكن الاستيراد إلى دراسة تحتوي بنوداً. صدّر الدراسة أولاً أو استخدم مشروعاً فارغاً لمنع التكرار.'
            ))

    def action_import_plan(self):
        self.ensure_one()
        values = self._validated_values()
        # Validation above occurs before this atomic multi-create; an ORM error also
        # rolls this savepoint back, so a failed import leaves no partial plan lines.
        with self.env.cr.savepoint():
            # The row lock makes simultaneous imports queue, then the second
            # request sees the lines created by the first request and stops.
            self._lock_empty_project_plan()
            paths = {}
            Line = self.env['abroj.cost.plan.line']
            for values_row in values:
                parent_path = values_row.pop('_parent_path')
                path = values_row.pop('_path')
                if parent_path:
                    values_row['parent_id'] = paths[parent_path]
                line = Line.create(values_row)
                if path:
                    paths[path] = line.id
        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': _('تم الاستيراد'),
                'message': _('تم إنشاء %(count)s بند دراسة للمشروع.') % {'count': len(values)},
                'type': 'success',
                'sticky': False,
                'next': {'type': 'ir.actions.act_window_close'},
            },
        }

    @staticmethod
    def _download_action(wizard, filename):
        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/?model=%s&id=%s&field=export_file&filename_field=export_filename&download=true' % (
                wizard._name, wizard.id,
            ),
            'target': 'self',
        }

    def _template_workbook(self):
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font
        except ImportError as error:
            raise UserError(_('لا تتوفر مكتبة إنشاء ملفات Excel في بيئة أودو.')) from error
        workbook = Workbook()
        plan_sheet = workbook.active
        plan_sheet.title = 'بنود الدراسة'
        plan_sheet.append(IMPORT_COLUMNS)
        plan_sheet.append((
            '', 'المطبخ', 'detailed', 0,
            0, 0, 0, 0, 0, '', 'المطبخ', 'section',
        ))
        category = self.env['abroj.cost.category'].search([
            ('company_id', '=', self.project_id.company_id.id), ('active', '=', True),
        ], limit=1)
        if category:
            plan_sheet.append((
                category.name, 'مثال بند', 'lump_sum', 1,
                0, 0, 0, 0, 1500, 'اكتب ملاحظتك هنا', 'المطبخ/مثال بند', 'item',
            ))
            # Keep the exact category name importable without allowing formulas.
            plan_sheet['A3'].data_type = 's'
        for cell in plan_sheet[1]:
            cell.font = Font(bold=True)
        for column, width in {'A': 24, 'B': 32, 'C': 20, 'D': 12, 'E': 20, 'F': 22, 'G': 20, 'H': 22, 'I': 18, 'J': 36, 'K': 48, 'L': 18}.items():
            plan_sheet.column_dimensions[column].width = width
        instructions = workbook.create_sheet('التعليمات')
        instructions.append((_('العمود'), _('الوصف')))
        for row in [
            ('category', _('اسم نوع عمل موجود في الشركة المختارة؛ مطلوب للبند المسعّر فقط. لن ينشئ النظام فئة أو مادة.')),
            ('name', _('اسم بند الدراسة.')),
            ('pricing_method', 'detailed / supply_only / labor_only / supply_install / lump_sum'),
            ('quantity', _('كمية موجبة، ومنزلتان عشريتان كحد أقصى.')),
            ('*_cost', _('سعر غير سالب ومنزلتان عشريتان كحد أقصى.')),
            ('notes', _('اختياري.')),
            ('path', _('اختياري للبند الجذري. للشجرة اكتب المسار كاملاً، وينتهي باسم هذا الصف، مثل المطبخ/سباكة/تمديدات.')),
            ('node_kind', _('item لبند مسعّر أو section لقسم تجميعي. القسم لا يحتاج نوع عمل أو سعر.')),
            (_('ترتيب الصفوف'), _('اكتب القسم الأب قبل أبنائه. ترتيب الصفوف يحدد ترتيب البنود داخل كل قسم.')),
            (_('الأمثلة'), _('عدّل أو احذف صفوف المثال قبل الاستيراد. إذا لم توجد فئات فعالة، أضف نوع عمل في النظام ثم أضف بنداً في الملف.')),
            (_('الدراسة الحالية'), _('الاستيراد إلى دراسة فارغة فقط؛ لا يحدث البنود الموجودة ولا يكررها.')),
        ]:
            instructions.append(row)
        for cell in instructions[1]:
            cell.font = Font(bold=True)
        instructions.column_dimensions['A'].width = 24
        instructions.column_dimensions['B'].width = 90
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    def action_download_template(self):
        self.ensure_one()
        self._ensure_project_scope()
        self.write({
            'export_file': base64.b64encode(self._template_workbook()),
            'export_filename': 'abroj_cost_plan_template.xlsx',
        })
        return self._download_action(self, self.export_filename)

    def action_export_plan(self):
        self.ensure_one()
        self._ensure_project_scope()
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font
        except ImportError as error:
            raise UserError(_('لا تتوفر مكتبة إنشاء ملفات Excel في بيئة أودو.')) from error
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'بنود الدراسة'
        sheet.append(IMPORT_COLUMNS)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        # Share the UI's authoritative depth-first order, including dragged sections.
        ordered_ids = [row['id'] for row in self.project_id.get_study_tree_data()]
        for line in self.env['abroj.cost.plan.line'].browse(ordered_ids):
            sheet.append((
                self._safe_excel_text(line.category_id.name if line.category_id else ''),
                self._safe_excel_text(line.name),
                line.pricing_method,
                line.quantity, line.material_unit_cost, line.auxiliary_unit_cost,
                line.labor_unit_cost, line.inclusive_unit_cost, line.lump_sum_cost,
                self._safe_excel_text(line.notes),
                self._safe_excel_text(self._plan_path(line)),
                line.node_kind,
            ))
        output = BytesIO()
        workbook.save(output)
        filename = 'abroj_cost_plan_%s.xlsx' % (self.project_id.name or self.project_id.id)
        self.write({
            'export_file': base64.b64encode(output.getvalue()),
            'export_filename': filename,
        })
        return self._download_action(self, filename)

    @staticmethod
    def _safe_excel_text(value):
        """Keep user-controlled values as literals when an XLSX is opened."""
        text = str(value or '')
        return "'%s" % text if text.startswith(('=', '+', '-', '@')) else text

    @staticmethod
    def _plan_path(line):
        names = []
        current = line
        while current:
            names.append(current.name)
            current = current.parent_id
        return '/'.join(reversed(names))
