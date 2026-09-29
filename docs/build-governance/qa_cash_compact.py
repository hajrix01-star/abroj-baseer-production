"""Compact cash-report presentation regression on committed QA company 9.

No ledger or configuration writes. Native report audit may persist via vendor
independent transactions; this script rolls back its own transaction.
"""
import json
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

assert env.cr.dbname == 'baseer_reports_qa_20260907'
c = env['res.company'].browse(9).exists()
assert c and c.name == 'QA الفئات النقدية'
scoped = env(context=dict(env.context, allowed_company_ids=c.ids))
report = scoped['eh.account.dynamic.report'].search([('code', '=', 'baseer_cash_categories')], limit=1)
assert report
root = Path('/mnt/qa-evidence')
previous = json.loads((root / 'cash_drilldown_checks.json').read_text(encoding='utf-8'))
source_baseline = {(item['mode'], item['line_id']): item for item in previous['action_evidence']}
checks, modes = [], []


def passed(condition, label, **evidence):
    assert condition, (label, evidence)
    checks.append({'check': label, 'passed': True, **evidence})


def money(actual, expected, label):
    passed(Decimal(str(actual)).quantize(Decimal('.01')) == Decimal(str(expected)).quantize(Decimal('.01')),
        label, actual=str(actual), expected=str(expected))


def is_number(line):
    value = line['columns'][0]['value']
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def assert_canonical(lines, include_tax, label):
    ids = [line['id'] for line in lines]
    passed(len(ids) == len(set(ids)), label + ': stable unique row IDs')
    passed(not any('/cash-account-' in key for key in ids), label + ': no redundant receipt account children')
    passed('baseer-total-receipts' not in ids and 'baseer-total-payments' not in ids,
        label + ': no repeated receipt/payment subtotal rows')
    passed(ids.count('receipts') == 1 and ids.count('payments') == 1, label + ': one header per cash direction')
    if include_tax:
        passed('baseer-total-displayed_net_movement' not in ids and ids.count('baseer-total-actual_net_movement') == 1,
            label + ': gross actual net appears once')
    else:
        passed(ids.count('baseer-total-displayed_net_movement') == 1 and ids.count('baseer-total-actual_net_movement') == 1,
            label + ': separate selected-basis net and actual cash reconciling amount')


try:
    for include_tax, mode in [(True, 'gross'), (False, 'tax-excluded')]:
        options = {'date': {'mode': 'range', 'date_from': '2026-09-01', 'date_to': '2026-09-30'},
            'company_ids': c.ids, 'posted_only': True, 'baseer_include_tax': include_tax,
            'lazy_expand': True, 'unfold_all': False}
        payload = report.render(options, use_cache=False)
        lines, totals = payload['lines'], payload['totals']
        for key, expected in {'receipts': 460 if include_tax else 400,
                'payments': -731 if include_tax else -677,
                'displayed_net_movement': -271 if include_tax else -277,
                'actual_net_movement': -271, 'opening_cash_balance': 1000,
                'closing_cash_balance': 729, 'balance_check': 0,
                'excluded_tax_bridge': 0 if include_tax else 6}.items():
            money(totals[key], expected, mode + ': unchanged ' + key)
        assert_canonical(lines, include_tax, mode + ' screen')
        by_id = {line['id']: line for line in lines}
        # Locate by the known opening/check children, independent of label language.
        opening = by_id['baseer-total-opening_cash_balance']
        audit = by_id[opening['parent_id']]
        passed(not is_number(audit) and audit.get('unfoldable') and audit.get('unfolded') is False,
            mode + ': audit parent is nonnumeric and collapsed initially')
        audit_children = [line for line in lines if line.get('parent_id') == audit['id']]
        required = ['baseer-total-opening_cash_balance', 'baseer-total-closing_cash_balance', 'baseer-total-balance_check']
        for key in required:
            child = by_id[key]
            passed(child.get('parent_id') == audit['id'] and child['level'] == audit['level'] + 1,
                mode + ': correct audit nesting ' + key)
        notes = [line for line in lines if line['id'].startswith('baseer-note-')]
        passed(bool(notes) and all(line.get('parent_id') == audit['id'] and line['level'] == audit['level'] + 1 for line in notes),
            mode + ': explanatory notes belong inside collapsed audit detail')
        if not include_tax:
            actual = by_id['baseer-total-actual_net_movement']
            passed(actual.get('parent_id') is not None, mode + ': actual cash amount is in detail rather than repeated at top level')
        for line in lines:
            parent_id = line.get('parent_id')
            if parent_id:
                passed(parent_id in by_id and line['level'] == by_id[parent_id]['level'] + 1,
                    mode + ': consistent hierarchy level ' + line['id'])

        numeric_count = 0
        for line in lines:
            if not is_number(line):
                continue
            numeric_count += 1
            action = report.get_drilldown_for_line(options, line['id'])
            passed(isinstance(action, dict) and action.get('type') == 'ir.actions.act_window'
                and action.get('res_model') in ('account.move', 'account.move.line'),
                mode + ': numeric amount retains native action ' + line['id'])
            model = scoped[action['res_model']]
            sources = model.browse(action['res_id']).exists() if action.get('res_id') else model.search(action['domain'])
            passed(all(source.company_id == c for source in sources), mode + ': source company scope ' + line['id'])
            originals = sources if sources._name == 'account.move' else sources.move_id
            passed(all(move.state == 'posted' for move in originals), mode + ': source posted scope ' + line['id'])
            baseline = source_baseline[(mode, line['id'])]
            passed(action['res_model'] == baseline['model'] and set(sources.ids) == set(baseline['source_ids']),
                mode + ': original source provenance unchanged ' + line['id'])
        passed(numeric_count > 10, mode + ': monetary drilldown coverage includes categories and audit amounts', count=numeric_count)
        zero = by_id['baseer-total-balance_check']
        money(zero['columns'][0]['value'], 0, mode + ': zero reconciliation still present in drillable audit detail')

        # Native PDF pipeline: compare real render rows and canonical payload,
        # including audit details in exports, without reintroducing duplicate rows.
        pdf_model = scoped['report.eh_account_dynamic_reports.report_dynamic_pdf_template']
        pdf_values = pdf_model._get_report_values(report.ids, data={'options': options})
        pdf = pdf_values['rendered'][0]
        assert_canonical(pdf['payload']['lines'], include_tax, mode + ' PDF')
        passed([line['id'] for line in pdf['lines']] == [line['id'] for line in pdf['payload']['lines']],
            mode + ': PDF output has exactly canonical rows without hidden duplicate totals', count=len(pdf['lines']))

        content = report.render_xlsx(options, use_cache=False)
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        header = next(index for index, row in enumerate(rows) if row[0] == payload['columns'][0]['name']
            and len(row) > 1 and row[1] == payload['columns'][1]['name'])
        data_rows = rows[header + 1:header + 1 + len(lines)]
        passed([row[0] for row in data_rows] == [line['name'] for line in lines],
            mode + ': XLSX output contains exactly canonical report rows', count=len(data_rows))
        for row, line in zip(data_rows, lines):
            if is_number(line):
                money(row[1], line['columns'][0]['value'], mode + ': XLSX monetary value ' + line['id'])
        remaining = rows[header + 1 + len(lines):]
        passed(not any(len(row) > 1 and isinstance(row[1], (int, float)) for row in remaining),
            mode + ': XLSX has no additional monetary footer duplicating totals')
        workbook.close()
        modes.append({'mode': mode, 'canonical_row_count': len(lines), 'numeric_action_count': numeric_count,
            'audit_child_count': len(audit_children), 'pdf_row_count': len(pdf['lines']), 'xlsx_data_row_count': len(data_rows)})

    env.cr.rollback()
    (root / 'cash_compact_checks.json').write_text(json.dumps({'company_id': 9, 'checks': checks, 'modes': modes,
        'ledger_config_mutations': False, 'audit_retention_note': 'Vendor durable failure audit may remain in QA.'},
        ensure_ascii=False, indent=2), encoding='utf-8')
    print('CASH_COMPACT_QA_SUCCESS', len(checks), 'checks;', json.dumps(modes))
except Exception:
    env.cr.rollback()
    raise
