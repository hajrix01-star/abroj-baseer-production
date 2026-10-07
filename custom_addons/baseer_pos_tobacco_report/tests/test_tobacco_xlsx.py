from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from time import monotonic
from urllib.parse import parse_qs, urlsplit
from unittest.mock import patch

from openpyxl import load_workbook

from odoo import Command, api
from odoo.exceptions import AccessError, UserError
from odoo.tests.common import HttpCase, TransactionCase, tagged

from ..models import tobacco_xlsx_export
from ..models.tobacco_report import BaseerPosTobaccoReport
from ..models.tobacco_report_wizard import TOBACCO_EXPORT_LOCK_NAMESPACE


def _options(company_id, **changes):
    options = {
        'company_id': company_id, 'period': 'month', 'month': 9, 'year': 2026,
        'date_from': '2026-09-01', 'date_to': '2026-09-30',
        'show_products': False, 'page': 1,
    }
    options.update(changes)
    return options


@tagged('post_install', '-at_install')
class TestTobaccoXlsx(TransactionCase):
    def setUp(self):
        super().setUp()
        self.env.company.currency_id = self.env.ref('base.SAR')
        self.wizards = self.env['baseer.pos.tobacco.report.wizard']
        self.exports = self.env['baseer.pos.tobacco.xlsx']
        self.options = _options(self.env.company.id)

    def _report(self, *, show_products=False):
        options = _options(self.env.company.id, show_products=show_products)
        company, values = self.wizards._hub_wizard_values(options)
        wizard = self.wizards.with_company(company).new(values)
        report = self.env['report.baseer_pos_tobacco_report.report_pos_tobacco_fees'].with_company(
            company,
        )._build_report(wizard, for_export=True)
        source = report['_export']
        sample = {
            'date': datetime(2026, 9, 1, 3, 0), 'order': '=SUM(1,2)',
            'products': '@Item × 1', 'type': 'Sale',
            'debit': Decimal('0.00'), 'credit': Decimal('12.30'),
            'running': Decimal('12.30'),
        }
        source['rows'] = [sample]
        source['exceptions'] = [{
            'date': sample['date'], 'order': '+ORDER',
            'computed_total': Decimal('17.31'), 'saved_total': Decimal('17.30'),
        }]
        source.update(debit_total=Decimal('0.00'), credit_total=Decimal('12.30'),
                      closing=Decimal('12.30'))
        report['rows'] = [{
            'date': '01-09-2026 03:00', 'order': sample['order'],
            'products': sample['products'], 'type': 'Sale',
            'debit': '0.00', 'credit': '12.30', 'running': '12.30',
        }]
        report['exceptions'] = [{
            'date': '01-09-2026 03:00', 'order': '+ORDER',
            'computed_total': '17.31', 'saved_total': '17.30',
        }]
        return report

    def test_export_raw_source_matches_screen_without_rpc_leak(self):
        currency = self.env.company.currency_id
        report_model = self.env['report.baseer_pos_tobacco_report.report_pos_tobacco_fees']

        class Product:
            def with_context(self, **kwargs):
                return self

            display_name = 'Shisha'

        class Line:
            qty = 1
            product_id = Product()

        class Order:
            currency_id = currency
            date_order = datetime(2026, 9, 1)
            name = 'ORDER-1'

        data = {'rows': [], 'exceptions': []}
        with patch.object(BaseerPosTobaccoReport, '_order_fee_data', return_value=(
            [(Line(), Decimal('12.30'))], True, Decimal('18.31'), Decimal('18.30'),
        )):
            rows, exceptions, debit, credit, closing = report_model._collect_rows(
                [Order()], currency, Decimal('0.01'), export_data=data,
            )
        self.assertEqual(data['rows'][0]['date'], datetime(2026, 9, 1, 3))
        self.assertEqual(data['rows'][0]['credit'], Decimal('12.30'))
        self.assertEqual(data['rows'][0]['running'], Decimal(rows[0]['running']))
        self.assertEqual(data['exceptions'][0]['computed_total'], Decimal('18.31'))
        self.assertEqual(data['exceptions'][0]['saved_total'], Decimal(exceptions[0]['saved_total']))
        self.assertEqual((debit, credit, closing),
                         (Decimal('0'), Decimal('12.30'), Decimal('12.30')))
        wizard = self.wizards.new({'company_id': self.env.company.id,
                                   'date_from': date(2026, 9, 1), 'date_to': date(2026, 9, 30)})
        ordinary = report_model._build_report(wizard)
        self.assertNotIn('_export', ordinary)
        hub = self.wizards.get_hub_report(self.options)
        self.assertNotIn('_export', hub)

    def test_numeric_dates_a4_exceptions_and_formula_safety(self):
        for products in (False, True):
            with self.subTest(products=products):
                report = self._report(show_products=products)
                payload = self.exports._render_workbook(report)
                self.assertLessEqual(len(payload), 5 * 1024 * 1024)
                book = load_workbook(BytesIO(payload))
                main, errors = book.worksheets
                self.assertEqual(len(book.worksheets), 2)
                self.assertEqual(main['A12'].value, datetime(2026, 9, 1, 3))
                self.assertEqual(main['B12'].value, '=SUM(1,2)')
                self.assertEqual(main['B12'].data_type, 's')
                self.assertTrue(main['B12'].alignment.wrap_text)
                self.assertEqual(main['B3'].alignment.horizontal, 'left')
                self.assertEqual(main['D12'].value if products else main['C12'].value,
                                 '@Item × 1' if products else 'Sale')
                incoming = main['F12' if products else 'E12']
                self.assertEqual(Decimal(str(incoming.value)), Decimal('12.30'))
                self.assertEqual(incoming.number_format, '#,##0.00')
                self.assertEqual(incoming.data_type, 'n')
                self.assertEqual(incoming.font.color.rgb[-6:], '243248')
                self.assertEqual(main['E12' if products else 'D12'].font.color.rgb[-6:],
                                 'AAB7C4')
                self.assertEqual(main['G12' if products else 'F12'].value, 12.3)
                self.assertGreaterEqual(main.row_dimensions[12].height, 25)
                self.assertEqual(errors['B7'].value, '+ORDER')
                self.assertEqual(errors['C7'].value, 17.31)
                self.assertEqual(errors['D7'].value, 17.3)
                self.assertEqual(str(main.page_setup.paperSize), main.PAPERSIZE_A4)
                self.assertEqual(main.page_setup.orientation,
                                 'landscape' if products else 'portrait')
                self.assertEqual(main.page_setup.fitToWidth, 1)
                self.assertEqual(main.print_title_rows, '$1:$10')
                self.assertIn('$A$1:', main.print_area)
                self.assertEqual(errors.print_title_rows, '$1:$6')
                self.assertIn('&P / &N', main.oddFooter.center.text)
                for sheet in book:
                    for row in sheet:
                        for cell in row:
                            self.assertNotEqual(cell.data_type, 'f')
                            self.assertIsNone(cell.hyperlink)

        refund = self._report()
        refund['_export']['rows'][0].update(debit=Decimal('12.30'), credit=Decimal('0.00'),
                                            running=Decimal('-12.30'))
        refund['_export'].update(debit_total=Decimal('12.30'), credit_total=Decimal('0.00'),
                                 closing=Decimal('-12.30'))
        refund_sheet = load_workbook(BytesIO(self.exports._render_workbook(refund))).active
        self.assertEqual(refund_sheet['D12'].font.color.rgb[-6:], 'D3262E')
        self.assertEqual(refund_sheet['E12'].font.color.rgb[-6:], 'AAB7C4')
        self.assertEqual(refund_sheet['F12'].font.color.rgb[-6:], 'D3262E')

    def test_arabic_rtl_empty_and_page_two_exports_full_period(self):
        report = self._report(show_products=True)
        report['is_rtl'] = True
        sheet = load_workbook(BytesIO(self.exports._render_workbook(report))).active
        self.assertTrue(sheet.sheet_view.rightToLeft)
        self.assertIn('رسوم التبغ', sheet['A1'].value)
        self.assertEqual(sheet['D12'].value, '@Item × 1')
        self.assertEqual(sheet['B3'].alignment.horizontal, 'right')

        report['rows'] = []
        report['exceptions'] = []
        report['_export'].update(rows=[], exceptions=[], debit_total=Decimal('0.00'),
                                 credit_total=Decimal('0.00'), closing=Decimal('0.00'))
        empty_book = load_workbook(BytesIO(self.exports._render_workbook(report)))
        self.assertEqual(empty_book.active['G12'].value, 0)
        self.assertEqual(empty_book.worksheets[1]['A7'].value, 'لا توجد فروق')

        full = self._report()
        with patch.object(BaseerPosTobaccoReport, '_build_report', return_value=full) as build:
            action = self.wizards.export_hub_xlsx(_options(self.env.company.id, page=2))
        self.assertTrue(build.call_args.kwargs['for_export'])
        export_id = int(urlsplit(action['url']).path.rsplit('/', 1)[1])
        self.assertEqual(parse_qs(urlsplit(action['url']).query)['company_id'],
                         [str(self.env.company.id)])
        stored = self.exports.browse(export_id)
        self.assertEqual(stored.show_products, False)
        import base64
        exported = load_workbook(BytesIO(base64.b64decode(stored.file_data)))
        self.assertEqual(exported.active['B12'].value, '=SUM(1,2)')
        self.assertEqual(exported.worksheets[1]['B7'].value, '+ORDER')

    def test_long_order_and_product_are_wrapped_without_clipping(self):
        report = self._report(show_products=True)
        report['_export']['rows'][0]['order'] = 'Long POS reference ' * 5
        report['_export']['rows'][0]['products'] = 'Long tobacco product name ' * 8
        sheet = load_workbook(BytesIO(self.exports._render_workbook(report))).active
        self.assertTrue(sheet['B12'].alignment.wrap_text)
        self.assertTrue(sheet['D12'].alignment.wrap_text)
        self.assertEqual(sheet['B12'].value, report['_export']['rows'][0]['order'])
        self.assertEqual(sheet['D12'].value, report['_export']['rows'][0]['products'])
        self.assertGreaterEqual(sheet.row_dimensions[12].height, 90)

    def test_precision_size_time_reject_before_transient(self):
        report = self._report()
        before = self.exports.search_count([])
        for number in (Decimal('1000000000000.00'), Decimal('0.001')):
            report['_export']['rows'][0]['running'] = number
            with self.assertRaises(UserError):
                self.exports._render_workbook(report)
        report['_export']['rows'][0]['running'] = Decimal('999999999999.99')
        book = load_workbook(BytesIO(self.exports._render_workbook(report)))
        self.assertEqual(Decimal(str(book.active['F12'].value)).quantize(Decimal('0.01')),
                         Decimal('999999999999.99'))
        report['_export']['exceptions'][0]['saved_total'] = Decimal('1000000000000.00')
        with self.assertRaises(UserError):
            self.exports._render_workbook(report)
        report['_export']['exceptions'][0]['saved_total'] = Decimal('17.30')
        with patch.object(tobacco_xlsx_export, '_FILE_LIMIT', 1):
            with self.assertRaises(UserError):
                self.exports._render_workbook(report)
        with self.assertRaises(UserError):
            self.exports._render_workbook(report, started_at=monotonic() - 11)
        self.assertEqual(self.exports.search_count([]), before)

    def test_company_lock_rejects_without_build_or_file_then_releases(self):
        before = self.exports.search_count([])
        with self.env.registry.cursor() as locker:
            locker.execute('SELECT pg_try_advisory_xact_lock(%s, %s)', (
                TOBACCO_EXPORT_LOCK_NAMESPACE, self.env.company.id,
            ))
            self.assertTrue(locker.fetchone()[0])
            with patch.object(BaseerPosTobaccoReport, '_build_report',
                              side_effect=AssertionError('POS must not be queried')) as build:
                with self.assertRaisesRegex(UserError, 'export is running'):
                    self.wizards.export_hub_xlsx(self.options)
            build.assert_not_called()
            self.assertEqual(self.exports.search_count([]), before)
            locker.rollback()
        action = self.wizards.export_hub_xlsx(self.options)
        self.assertEqual(action['target'], 'download')
        self.assertEqual(self.exports.search_count([]), before + 1)

    def test_lock_is_per_company_and_released_after_commit(self):
        other = self.env['res.company'].create({
            'name': 'Independent tobacco XLSX company',
            'currency_id': self.env.ref('base.SAR').id,
        })
        self.env.user.write({'company_ids': [Command.link(other.id)]})
        with self.env.registry.cursor() as locker:
            locker.execute('SELECT pg_try_advisory_xact_lock(%s, %s)', (
                TOBACCO_EXPORT_LOCK_NAMESPACE, self.env.company.id,
            ))
            self.assertTrue(locker.fetchone()[0])
            other_model = self.wizards.with_company(other).with_context(
                allowed_company_ids=[other.id],
            )
            action = other_model.export_hub_xlsx(_options(other.id))
            self.assertEqual(action['target'], 'download')
            locker.commit()
        # A different DB transaction now acquires the released first-company key.
        with self.env.registry.cursor() as verifier:
            verifier.execute('SELECT pg_try_advisory_xact_lock(%s, %s)', (
                TOBACCO_EXPORT_LOCK_NAMESPACE, self.env.company.id,
            ))
            self.assertTrue(verifier.fetchone()[0])
            verifier.rollback()

    def test_failed_export_releases_lock_on_request_rollback(self):
        with self.env.registry.cursor() as request_cursor:
            request_env = api.Environment(request_cursor, self.env.uid, {
                'allowed_company_ids': [self.env.company.id],
            })
            with patch.object(BaseerPosTobaccoReport, '_build_report',
                              side_effect=UserError('forced generation failure')):
                with self.assertRaisesRegex(UserError, 'forced generation failure'):
                    request_env['baseer.pos.tobacco.report.wizard'].export_hub_xlsx(
                        self.options,
                    )
            request_cursor.rollback()
        with self.env.registry.cursor() as verifier:
            verifier.execute('SELECT pg_try_advisory_xact_lock(%s, %s)', (
                TOBACCO_EXPORT_LOCK_NAMESPACE, self.env.company.id,
            ))
            self.assertTrue(verifier.fetchone()[0])
            verifier.rollback()

    def test_owner_company_and_accounting_pos_access(self):
        other_company = self.env['res.company'].create({'name': 'Other tobacco XLSX company'})
        account = self.env.ref('account.group_account_readonly')
        pos = self.env.ref('point_of_sale.group_pos_user')
        common = self.env.ref('base.group_user')
        users = []
        for name in ('owner', 'other'):
            users.append(self.env['res.users'].create({
                'name': 'Tobacco XLSX ' + name, 'login': 'tobacco.xlsx.' + name,
                'company_id': self.env.company.id,
                'company_ids': [Command.set((self.env.company | other_company).ids)],
                'group_ids': [Command.set((account | pos | common).ids)],
            }))
        owner, other = users
        export = self.exports.with_user(owner).create({
            'company_id': self.env.company.id, 'period': 'month', 'month': 9, 'year': 2026,
            'date_from': date(2026, 9, 1), 'date_to': date(2026, 9, 30),
            'file_data': 'ZGF0YQ==', 'file_name': 'safe.xlsx',
        })
        scoped = export.with_user(owner).with_context(allowed_company_ids=[self.env.company.id])
        self.assertEqual(scoped._download_record(export.id, self.env.company.id), export)
        with self.assertRaises(AccessError):
            export.with_user(other)._download_record(export.id, self.env.company.id)
        with self.assertRaises(AccessError):
            export.with_user(owner).with_context(allowed_company_ids=[other_company.id])._download_record(
                export.id, self.env.company.id)
        with self.assertRaises(AccessError):
            scoped._download_record(export.id, other_company.id)
        owner.group_ids = [Command.set((account | common).ids)]
        with self.assertRaises(AccessError):
            scoped._download_record(export.id, self.env.company.id)
        with self.assertRaises(AccessError):
            scoped.read(['file_data'])
        owner.group_ids = [Command.set((pos | common).ids)]
        with self.assertRaises(AccessError):
            scoped._download_record(export.id, self.env.company.id)
        with self.assertRaises(AccessError):
            scoped.read(['file_data'])


@tagged('post_install', '-at_install')
class TestTobaccoXlsxHttp(HttpCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.company = cls.env.company
        cls.company.currency_id = cls.env.ref('base.SAR')
        cls.other_company = cls.env['res.company'].create({'name': 'Other tobacco HTTP company'})
        groups = (cls.env.ref('base.group_user') |
                  cls.env.ref('account.group_account_readonly') |
                  cls.env.ref('point_of_sale.group_pos_user'))
        users = []
        for name in ('owner', 'other'):
            users.append(cls.env['res.users'].create({
                'name': 'Tobacco XLSX HTTP ' + name, 'login': 'tobacco.xlsx.http.' + name,
                'password': 'tobacco-xlsx-http-test', 'company_id': cls.company.id,
                'company_ids': [Command.set((cls.company | cls.other_company).ids)],
                'group_ids': [Command.set(groups.ids)],
            }))
        cls.owner, cls.other = users

    def test_private_route_and_generic_content(self):
        owner_model = self.env['baseer.pos.tobacco.report.wizard'].with_user(
            self.owner,
        ).with_context(allowed_company_ids=[self.company.id])
        action = owner_model.export_hub_xlsx(_options(self.company.id))
        record_id = int(urlsplit(action['url']).path.rsplit('/', 1)[1])
        generic = '/web/content?model=baseer.pos.tobacco.xlsx&id=%s&field=file_data&download=true' % record_id
        self.authenticate(self.owner.login, 'tobacco-xlsx-http-test')
        owned = self.url_open(action['url'])
        self.assertEqual(owned.status_code, 200)
        self.assertIn('spreadsheetml.sheet', owned.headers['Content-Type'])
        self.assertIn('no-store', owned.headers['Cache-Control'])
        self.assertEqual(len(load_workbook(BytesIO(owned.content)).worksheets), 2)
        self.assertEqual(self.url_open(action['url'].replace(
            'company_id=%s' % self.company.id,
            'company_id=%s' % self.other_company.id,
        )).status_code, 404)
        self.assertEqual(self.url_open('/baseer/pos/tobacco/export/2147483647?company_id=%s' %
                                       self.company.id).status_code, 404)
        self.assertNotEqual(self.url_open(generic).status_code, 200)
        self.authenticate(self.other.login, 'tobacco-xlsx-http-test')
        self.assertEqual(self.url_open(action['url']).status_code, 404)
        self.assertNotEqual(self.url_open(generic).status_code, 200)
        self.authenticate(self.owner.login, 'tobacco-xlsx-http-test')
        self.owner.company_id = self.other_company
        self.assertEqual(self.url_open(action['url']).status_code, 404)
        self.opener.cookies.clear()
        self.assertNotEqual(self.url_open(action['url'], allow_redirects=False).status_code, 200)
