"""Bounded, pure presentation of authoritative Baseer report documents.

No financial calculation, ORM access, URL resolution, or shared mutable state.
Missing values are blank; zero is a number. Invalid precision fails closed.
"""

from datetime import date
from base64 import b64encode
from decimal import Decimal
from html import escape
from io import BytesIO
from math import ceil
from pathlib import Path
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


MAX_ROWS = 5000
MAX_CELLS = 50000
MAX_OUTPUT_BYTES = 10 * 1024 * 1024
_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")
_KEY = re.compile(r"[A-Za-z][A-Za-z0-9_]{0,63}\Z")
_NUMERIC_METADATA = re.compile(r"[0-9\s:./\-–—+TZ]+\Z")
_INK = "243248"
_SHADE = "EDF1F4"
_LINE = "CCD3DA"
_FONT_FAMILY = "IBM Plex Sans Arabic"


class DocumentValidationError(ValueError):
    """The document cannot be exported completely and faithfully."""


def _require(condition, message):
    if not condition:
        raise DocumentValidationError(message)


def _text(value, limit, name, nonempty=False):
    _require(isinstance(value, str), f"{name} must be text")
    _require(len(value) <= limit and not _CONTROL.search(value), f"Invalid or oversized {name}")
    _require(not nonempty or bool(value.strip()), f"{name} must not be empty")


def _fields(value, allowed, required, name):
    _require(isinstance(value, dict), f"{name} must be a mapping")
    _require(required <= value.keys() <= allowed, f"Invalid {name} fields")


def _number(value, kind, decimals):
    _require(isinstance(value, Decimal) and value.is_finite(), "Numbers must be finite Decimal values")
    # Inspect tuple directly: normalize() would round under the ambient context.
    sign, digits, exponent = value.as_tuple()
    _require(len(digits) <= 64 and -64 <= exponent <= 64, "Numeric representation exceeds export limits")
    digits = list(digits)
    while len(digits) > 1 and digits[-1] == 0:
        digits.pop()
        exponent += 1
    significant = len(digits) if value else 1
    _require(significant <= 15, "Excel supports at most 15 significant digits")
    fractional = max(-exponent, 0) if value else 0
    integer = max(value.adjusted() + 1, 0) if value else 1
    _require(integer <= 15 and fractional <= 15, "Numeric magnitude or fractional precision exceeds export limits")
    if kind == "money":
        _require(fractional <= decimals, "Money exceeds declared currency precision")
        _require(integer + decimals <= 15, "Currency precision exceeds Excel's 15-digit budget")


def validate_document(doc):
    """Validate without mutation; return the original document on success."""
    names = {"title", "company", "currency", "lang", "decimals", "metadata", "columns", "rows", "notes"}
    _fields(doc, names, names, "document")
    for name in ("title", "company"):
        _text(doc[name], 256, name, nonempty=True)
    _text(doc["currency"], 16, "currency", nonempty=True)
    _require(doc["lang"] in ("ar", "en"), "Unsupported language")
    _require(type(doc["decimals"]) is int and 0 <= doc["decimals"] <= 2, "Currency decimals must be 0..2")
    _require(isinstance(doc["metadata"], list) and len(doc["metadata"]) <= 32, "Invalid metadata")
    for pair in doc["metadata"]:
        _require(isinstance(pair, (tuple, list)) and len(pair) == 2, "Metadata must contain label/value pairs")
        _text(pair[0], 128, "metadata label", nonempty=True)
        if type(pair[1]) is not date:
            _text(pair[1], 4096, "metadata value")
    _require(isinstance(doc["notes"], list) and len(doc["notes"]) <= 32, "Invalid notes")
    for note in doc["notes"]:
        _text(note, 2048, "note")
    columns = doc["columns"]
    _require(isinstance(columns, list) and 1 <= len(columns) <= 32, "Expected 1..32 columns")
    keys = set()
    for column in columns:
        _fields(column, {"key", "label", "type", "width"}, {"key", "label", "type"}, "column")
        _text(column["key"], 64, "column key", nonempty=True)
        _require(bool(_KEY.fullmatch(column["key"])) and column["key"] not in keys, "Invalid or duplicate column key")
        keys.add(column["key"])
        _text(column["label"], 256, "column label", nonempty=True)
        _require(column["type"] in ("text", "money", "percent", "date"), "Invalid column type")
        if "width" in column:
            _require(type(column["width"]) is int and 6 <= column["width"] <= 80, "Column width must be an integer in 6..80")
    rows = doc["rows"]
    _require(isinstance(rows, list) and len(rows) <= MAX_ROWS, "Too many rows; narrow the report period")
    _require(len(rows) * len(columns) <= MAX_CELLS, "Too many cells; narrow the report scope")
    text_bytes = 0
    for row in rows:
        _fields(row, {"cells", "kind", "indent"}, {"cells", "kind"}, "row")
        _require(row["kind"] in ("detail", "section", "total"), "Invalid row kind")
        _require(type(row.get("indent", 0)) is int and 0 <= row.get("indent", 0) <= 3, "Invalid row indent")
        _require(isinstance(row["cells"], dict) and row["cells"].keys() <= keys, "Invalid row cells")
        for column in columns:
            value = row["cells"].get(column["key"])
            if value is None:
                continue
            kind = column["type"]
            if kind == "text":
                _text(value, 4096, "cell text")
                text_bytes += len(value.encode("utf-8"))
                _require(text_bytes <= MAX_OUTPUT_BYTES, "Document text exceeds the 10 MiB limit; narrow the report scope")
            elif kind == "date":
                _require(type(value) is date, "Date cells require date (not datetime) values")
            else:
                _number(value, kind, doc["decimals"])
    return doc


class _BoundedBuffer(BytesIO):
    def write(self, data):
        _require(self.tell() + len(data) <= MAX_OUTPUT_BYTES, "Export exceeds the 10 MiB limit; narrow the report scope")
        return super().write(data)


class _HtmlParts(list):
    """Enforce the output budget while building, not after a huge allocation."""

    def __init__(self, values):
        super().__init__()
        self.byte_count = 0
        for value in values:
            self.append(value)

    def append(self, value):
        self.byte_count += len(value.encode("utf-8"))
        _require(self.byte_count <= MAX_OUTPUT_BYTES, "Export exceeds the 10 MiB limit; narrow the report scope")
        super().append(value)


def _literal(cell, value):
    cell.value = value
    if isinstance(value, str):
        # Explicit type prevents '=...' being interpreted as an Excel formula.
        cell.data_type = "s"


def _width(column):
    return column.get("width", 42 if column["type"] == "text" else 20)


def _height(text, width, size=10):
    """Conservative wrap estimate; never quietly clip an oversized Excel row."""
    characters = max(1, int(width * 0.75 * 10 / size))
    lines = sum(max(1, ceil(len(line) / characters)) for line in text.split("\n"))
    height = max(20, lines * (size + 4) + 6)
    _require(height <= 409, "Text cannot fit a readable Excel row; shorten the field or increase column width")
    return height


def render_xlsx(doc):
    """Return a self-contained workbook, with typed values and no active links."""
    validate_document(doc)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "التقرير" if doc["lang"] == "ar" else "Report"
    sheet.sheet_view.rightToLeft = doc["lang"] == "ar"
    sheet.sheet_view.showGridLines = False
    count = len(doc["columns"])
    align = "right" if doc["lang"] == "ar" else "left"

    def banner(row, value, bold=False, size=10):
        cell = sheet.cell(row, 1)
        _literal(cell, value)
        if count > 1:
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=count)
        cell.font = Font(name=_FONT_FAMILY, size=size, bold=bold, color=_INK)
        cell.alignment = Alignment(horizontal=align, vertical="top", wrap_text=True)
        sheet.row_dimensions[row].height = _height(value, sum(_width(c) for c in doc["columns"]), size)

    banner(1, doc["title"], True, 15)
    banner(2, doc["company"], True)
    banner(3, ("العملة: " if doc["lang"] == "ar" else "Currency: ") + doc["currency"])
    cursor = 4
    for label, value in doc["metadata"]:
        if type(value) is date:
            # Keep dates typed, including single-column documents.
            banner(cursor, label, True)
            cursor += 1
            cell = sheet.cell(cursor, 1, value)
            cell.number_format = "yyyy-mm-dd"
            cell.font = Font(name=_FONT_FAMILY, size=10, color=_INK)
            cell.alignment = Alignment(horizontal=align)
        else:
            banner(cursor, f"{label}: {value}")
        cursor += 1
    cursor += 1
    header = cursor
    for index, column in enumerate(doc["columns"], 1):
        cell = sheet.cell(header, index)
        _literal(cell, column["label"])
        cell.font = Font(name=_FONT_FAMILY, size=10, bold=True, color=_INK)
        cell.fill = PatternFill("solid", fgColor=_SHADE)
        cell.alignment = Alignment(horizontal=align, wrap_text=True, vertical="center")
        sheet.column_dimensions[get_column_letter(index)].width = _width(column)
    sheet.row_dimensions[header].height = max(_height(c["label"], _width(c)) for c in doc["columns"])
    money_format = "#,##0" + ("." + "0" * doc["decimals"] if doc["decimals"] else "")
    for row in doc["rows"]:
        cursor += 1
        sheet.row_dimensions[cursor].height = max(
            [_height(row["cells"].get(c["key"]) or "", _width(c) - row.get("indent", 0) * 2)
             for c in doc["columns"] if c["type"] == "text"] or [20])
        for index, column in enumerate(doc["columns"], 1):
            cell = sheet.cell(cursor, index)
            _literal(cell, row["cells"].get(column["key"]))
            numeric = column["type"] in ("money", "percent")
            ltr = numeric or column["type"] == "date" or (
                isinstance(cell.value, str) and bool(_NUMERIC_METADATA.fullmatch(cell.value)))
            cell.font = Font(name=_FONT_FAMILY, size=10, bold=row["kind"] != "detail", color=_INK)
            cell.alignment = Alignment(horizontal="right" if numeric else align, vertical="top", wrap_text=not numeric,
                                       indent=row.get("indent", 0) if column["type"] == "text" else 0,
                                       readingOrder=1 if ltr else (2 if doc["lang"] == "ar" else 1))
            cell.border = Border(bottom=Side(style="thin", color=_LINE))
            if row["kind"] != "detail":
                cell.fill = PatternFill("solid", fgColor=_SHADE)
            if column["type"] == "money":
                cell.number_format = money_format
            elif column["type"] == "percent":
                cell.number_format = "0.00" + "#" * 13 + "%"
            elif column["type"] == "date":
                cell.number_format = "yyyy-mm-dd"
    for note in doc["notes"]:
        cursor += 1
        banner(cursor, note)
    sheet.freeze_panes = f"A{header + 1}"
    sheet.print_title_rows = f"{header}:{header}"
    sheet.print_area = f"A1:{get_column_letter(count)}{cursor}"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape" if sum(_width(c) for c in doc["columns"]) > 95 else "portrait"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.oddFooter.center.text = "&P / &N"
    sheet.oddFooter.center.font = _FONT_FAMILY
    sheet.oddFooter.center.size = 9
    stream = _BoundedBuffer()
    workbook.save(stream)
    return stream.getvalue()


def _display(value, kind, decimals):
    if value is None:
        return ""
    if type(value) is date:
        return value.isoformat()
    if kind == "money":
        return format(value, f",.{decimals}f")
    if kind == "percent":
        sign, digits, exponent = value.as_tuple()
        scaled = Decimal((sign, digits, exponent + 2))
        return format(scaled, "f").rstrip("0").rstrip(".") + "%" if scaled and scaled.as_tuple().exponent < 0 else format(scaled, "f") + "%"
    return value


def _font_styles():
    """Embed only the two licensed package assets; never resolve client URLs."""
    folder = Path(__file__).resolve().parent / "static" / "fonts"
    styles = []
    for filename, weight in (("IBMPlexSansArabic-Regular.ttf", 400),
                             ("IBMPlexSansArabic-Bold.ttf", 700)):
        try:
            with (folder / filename).open("rb") as source:
                content = source.read(1024 * 1024 + 1)
        except OSError as error:
            raise DocumentValidationError("The bundled report font is unavailable") from error
        _require(0 < len(content) <= 1024 * 1024, "The bundled report font exceeds its size limit")
        encoded = b64encode(content).decode("ascii")
        styles.append(f"@font-face{{font-family:'{_FONT_FAMILY}';font-style:normal;font-weight:{weight};"
                      f"src:url(data:font/ttf;base64,{encoded}) format('truetype');}}")
    return "".join(styles)


def render_html(doc):
    """Return escaped HTML with embedded package fonts and no external resources."""
    validate_document(doc)
    rtl = doc["lang"] == "ar"
    direction = "rtl" if rtl else "ltr"
    align = "right" if rtl else "left"
    parts = _HtmlParts([f'<!doctype html><html lang="{doc["lang"]}" dir="{direction}"><head><meta charset="utf-8">',
             "<style>" + _font_styles(),
             f"body{{font-family:'{_FONT_FAMILY}',Arial,sans-serif;font-size:10pt;color:#243248;margin:0;}}"
             "h1{font-size:15pt;margin:0 0 5px;}p{margin:4px 0;}header{border-bottom:1px solid #CCD3DA;padding-bottom:10px;margin-bottom:12px;}"
             ".metadata,.notes{font-size:9pt;word-wrap:break-word;white-space:pre-wrap;}"
             ".metadata-value-ltr{direction:ltr;unicode-bidi:embed;display:inline-block;}"
             "table{border-collapse:collapse;width:100%;table-layout:fixed;}thead{display:table-header-group;}"
             "tr{page-break-inside:avoid;}td,th{padding:6px 8px;border-bottom:1px solid #CCD3DA;vertical-align:top;word-wrap:break-word;white-space:pre-wrap;}"
             f"th,td{{text-align:{align};}}"
             "th,.section,.total{background:#EDF1F4;font-weight:bold;}.section{page-break-after:avoid;}"
             ".before-total{page-break-after:avoid;}.total{page-break-before:avoid;}"
             ".numeric{direction:ltr;unicode-bidi:embed;text-align:right;white-space:nowrap;}"
             ".date{direction:ltr;unicode-bidi:embed;}.notes{margin-top:12px;}"
             "</style></head><body><header>", f"<h1>{escape(doc['title'])}</h1>",
             f"<p>{escape(doc['company'])}</p><p>{escape(doc['currency'])}</p>"])
    for label, value in doc["metadata"]:
        text = value.isoformat() if type(value) is date else value
        ltr = type(value) is date or bool(_NUMERIC_METADATA.fullmatch(text))
        value_span = (f'<span class="metadata-value-ltr" dir="ltr">{escape(text)}</span>' if ltr
                      else f'<span dir="auto">{escape(text)}</span>')
        parts.append(f'<p class="metadata"><strong>{escape(label)}:</strong> {value_span}</p>')
    parts.append("</header><table><colgroup>")
    widths = [_width(c) for c in doc["columns"]]
    total_width = sum(widths)
    for width in widths:
        parts.append(f'<col style="width:{width * 100 // total_width}%">')
    parts.append("</colgroup><thead><tr>")
    for column in doc["columns"]:
        parts.append(f"<th>{escape(column['label'])}</th>")
    parts.append("</tr></thead><tbody>")
    for index, row in enumerate(doc["rows"]):
        before = index + 1 < len(doc["rows"]) and doc["rows"][index + 1]["kind"] == "total"
        classes = row["kind"] + (" before-total" if before else "")
        parts.append(f'<tr class="{classes}">')
        for column in doc["columns"]:
            kind = column["type"]
            css = "numeric" if kind in ("money", "percent") else "date" if kind == "date" else "text"
            indent = row.get("indent", 0) * 12 + 8 if kind == "text" else 8
            padding = "right" if rtl else "left"
            value = _display(row["cells"].get(column["key"]), kind, doc["decimals"])
            if kind == "text" and _NUMERIC_METADATA.fullmatch(value):
                css = "date"
            parts.append(f'<td class="{css}" style="padding-{padding}:{indent}px">{escape(value)}</td>')
        parts.append("</tr>")
    parts.append('</tbody></table><div class="notes">')
    for note in doc["notes"]:
        parts.append(f"<p>{escape(note)}</p>")
    parts.append("</div></body></html>")
    return "".join(parts)
