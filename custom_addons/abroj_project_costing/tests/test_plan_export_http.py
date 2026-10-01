from io import BytesIO
from urllib.parse import parse_qs, urlsplit

from openpyxl import load_workbook

from odoo import Command
from odoo.tests.common import HttpCase, tagged


@tagged('post_install', '-at_install')
class TestAbrojPlanExportHttp(HttpCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.default_company = cls.env['res.company'].create({'name': 'Export HTTP default'})
        cls.company = cls.env['res.company'].create({
            'name': 'Export HTTP project', 'abroj_project_costing_enabled': True,
        })
        cls.foreign_company = cls.env['res.company'].create({
            'name': 'Export HTTP not permitted', 'abroj_project_costing_enabled': True,
        })
        group = cls.env.ref('abroj_project_costing.group_abroj_cost_user')
        cls.users = []
        for role in ('owner', 'other'):
            cls.users.append(cls.env['res.users'].with_context(no_reset_password=True).create({
                'name': 'Export HTTP %s' % role, 'login': 'abroj.export.http.%s' % role,
                'password': 'export-http-test', 'company_id': cls.default_company.id,
                'company_ids': [Command.set((cls.default_company | cls.company).ids)],
                'group_ids': [Command.set(group.ids)],
            }))
        cls.owner, cls.other = cls.users
        cls.owner_env = cls.env(user=cls.owner, context=dict(
            cls.env.context, allowed_company_ids=cls.company.ids,
        ))
        cls.category = cls.env['abroj.cost.category'].with_company(cls.company).create({
            'name': 'Export HTTP category', 'company_id': cls.company.id,
        })
        cls.project = cls.owner_env['abroj.cost.project'].create({'name': 'Export HTTP project'})
        cls.line = cls.owner_env['abroj.cost.plan.line'].create({
            'project_id': cls.project.id, 'category_id': cls.category.id,
            'name': 'HTTP study item', 'quantity': 3,
            'pricing_method': 'lump_sum', 'lump_sum_cost': 125,
        })

    def _export(self):
        action = self.project.action_export_plan()
        url = urlsplit(action['url'])
        wizard = self.owner_env['abroj.cost.plan.import.wizard'].browse(int(url.path.rsplit('/', 1)[1]))
        self.assertEqual(action['target'], 'download')
        self.assertEqual(parse_qs(url.query)['company_id'], [str(self.company.id)])
        return wizard, action['url']

    def _authenticate_owner(self):
        self.authenticate(self.owner.login, 'export-http-test')

    def test_export_preserves_company_context_without_relaxing_native_access(self):
        wizard, url = self._export()
        self._authenticate_owner()
        legacy = '/web/content?model=%s&id=%s&field=export_file&download=true' % (wizard._name, wizard.id)
        self.assertEqual(self.url_open(legacy).status_code, 404)
        response = self.url_open(url)
        self.assertEqual(response.status_code, 200)
        self.assertIn('attachment', response.headers['Content-Disposition'])
        self.assertIn('spreadsheetml.sheet', response.headers['Content-Type'])
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(workbook.active['B2'].value, 'HTTP study item')
        self.assertEqual(self.line.quantity, 3)
        self.assertEqual(self.project.estimated_total, 125)
        # Caller-supplied field/model never switches the content being served.
        self.assertEqual(self.url_open(url + '&field=import_file&model=res.users').content, response.content)

    def test_template_uses_same_authenticated_company_download(self):
        wizard = self.owner_env['abroj.cost.plan.import.wizard'].create({
            'project_id': self.project.id, 'company_id': self.company.id,
        })
        action = wizard.action_download_template()
        self._authenticate_owner()
        response = self.url_open(action['url'])
        self.assertEqual(response.status_code, 200)
        self.assertIn('abroj_cost_plan_template.xlsx', response.headers['Content-Disposition'])
        self.assertEqual(load_workbook(BytesIO(response.content)).active.max_column, 12)

    def test_download_rejects_other_user_company_tampering_and_missing_file(self):
        wizard, url = self._export()
        self.authenticate(self.other.login, 'export-http-test')
        self.assertEqual(self.url_open(url).status_code, 404)
        self._authenticate_owner()
        base = '/abroj/costing/plan/export/%s?company_id=' % wizard.id
        for company_id in (self.foreign_company.id, self.default_company.id, 'invalid', ''):
            with self.subTest(company_id=company_id):
                self.assertEqual(self.url_open(base + str(company_id)).status_code, 404)
        self.assertEqual(self.url_open('/abroj/costing/plan/export/2147483647?company_id=%s' % self.company.id).status_code, 404)
        empty = self.owner_env['abroj.cost.plan.import.wizard'].create({
            'project_id': self.project.id, 'company_id': self.company.id,
        })
        self.assertEqual(self.url_open(empty._download_action(empty, 'empty.xlsx')['url']).status_code, 404)
        self.company.abroj_project_costing_enabled = False
        self.assertEqual(self.url_open(url).status_code, 404)
        self.opener.cookies.clear()
        self.assertNotEqual(self.url_open(url, allow_redirects=False).status_code, 200)
