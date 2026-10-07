from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from time import monotonic
from types import SimpleNamespace
from urllib.parse import parse_qs, urlsplit
from unittest.mock import patch

from openpyxl import load_workbook

from odoo import Command
from odoo.exceptions import AccessError, UserError
from odoo.fields import Domain
from odoo.tests.common import HttpCase, TransactionCase, tagged

from ..models import vat_xlsx_export
from ..models.tax_report_wizard import BaseerTaxReportWizard


def _options(company_id, journal_ids=None):
    return {
        'company_id': company_id, 'period_type': 'quarter', 'year': 2026,
        'month': '9', 'quarter': '3', 'display_mode': 'simple',
        'journal_ids': journal_ids or [],
    }


@tagged('post_install', '-at_install')
class TestVatXlsx(TransactionCase):
    def setUp(self):
        super().setUp()
        self.company = self.env.company
        self.company.account_fiscal_country_id = self.env.ref('base.sa')
        self.model = self.env['baseer.tax.report.wizard']
        self.options = _options(self.company.id)

    def _workbook(self, report=None):
        report = report or self.model._hub_wizard(self.options)._build_report()
        payload = self.env['baseer.tax.report.xlsx']._render_workbook(report)
        self.assertLessEqual(len(payload), 1024 * 1024)
        return load_workbook(BytesIO(payload)).active

    def test_complete_return_matches_authoritative_decimals_and_a4_settings(self):
        report = self.model._hub_wizard(self.options)._build_report()
        sheet = self._workbook(report)
        self.assertEqual(len(sheet.parent.sheetnames), 1)
        self.assertEqual(sheet.max_row, 30)
        self.assertEqual([str(sheet[f'A{index}'].value) for index in range(10, 26)],
                         [str(index) for index in range(1, 17)])
        self.assertEqual(sheet['B10'].value, report['rows'][0]['name'].split('. ', 1)[-1])
        self.assertFalse(sheet['B10'].value.startswith('1. '))
        for index, row in enumerate(report['rows'], 10):
            for column, key in (('C', 'base'), ('D', 'tax')):
                expected = row[key]
                actual = sheet[f'{column}{index}'].value
                self.assertEqual(actual is None, expected is None)
                if expected is not None:
                    self.assertEqual(Decimal(str(actual)).quantize(Decimal('0.01')), expected)
                    self.assertEqual(sheet[f'{column}{index}'].data_type, 'n')
        self.assertIsNone(sheet['C22'].value)  # Box 13 has no statutory base.
        self.assertEqual(sheet['D10'].value, 0)
        self.assertEqual(sheet['D10'].font.color.rgb[-6:], 'AAB7C4')
        self.assertEqual(sheet['B3'].value.date() if isinstance(sheet['B3'].value, datetime)
                         else sheet['B3'].value, date(2026, 7, 1))
        self.assertEqual(str(sheet.page_setup.paperSize), sheet.PAPERSIZE_A4)
        self.assertEqual(sheet.page_setup.orientation, 'portrait')
        self.assertEqual(sheet.page_setup.fitToWidth, 1)
        self.assertEqual(sheet.print_title_rows, '$1:$9')
        self.assertIn('$A$1:$D$30', sheet.print_area)
        self.assertEqual(sheet.freeze_panes, 'C10')
        self.assertFalse(sheet.sheet_view.rightToLeft)
        self.assertEqual(sheet['B3'].alignment.horizontal, 'left')
        self.assertEqual(sheet['B5'].alignment.horizontal, 'left')
        self.assertGreaterEqual(sheet.column_dimensions['A'].width, 18)
        self.assertGreaterEqual(sheet.row_dimensions[18].height, 34)

    def test_arabic_journal_warning_exceptions_and_literal_text(self):
        journal = self.env['account.journal'].search([('company_id', '=', self.company.id)], limit=1)
        self.assertTrue(journal)
        options = _options(self.company.id, [journal.id])
        report = self.model.with_context(lang='ar_001')._hub_wizard(options)._build_report()
        report['journal_names'] = '=HYPERLINK("https://example.test", "X")'
        report['rows'][0]['name'] = '+SUM(1,1)'
        report['rows'][1]['name'] = '-1+2'
        report['rows'][2]['name'] = '@ATTACK'
        report['rows'][3]['name'] = '\u200b=SUM(1,1)'
        report['rows'][13]['tax'] = Decimal('-2.50')
        report['exception']['count'] = 3
        report['exception']['amount'] = Decimal('-7.17')
        sheet = self._workbook(report)
        self.assertTrue(sheet.sheet_view.rightToLeft)
        self.assertEqual(sheet['B3'].alignment.horizontal, 'right')
        self.assertEqual(sheet['B5'].alignment.horizontal, 'right')
        self.assertEqual(sheet['B11'].value, report['rows'][1]['name'].split('. ', 1)[-1])
        self.assertIn('ليس إجمالي الإقرار', sheet['A7'].value)
        self.assertEqual(sheet['B6'].value, report['journal_names'])
        self.assertEqual(sheet['C28'].value, 3)
        self.assertEqual(Decimal(str(sheet['D28'].value)), Decimal('-7.17'))
        self.assertEqual(sheet['D23'].font.color.rgb[-6:], 'D3262E')
        for row in sheet:
            for cell in row:
                self.assertNotEqual(cell.data_type, 'f')
                self.assertIsNone(cell.hyperlink)
        self.assertEqual(sheet['B10'].value, '+SUM(1,1)')

    def test_arabic_box_label_is_not_prefixed_twice(self):
        report = self.model.with_context(lang='ar_001')._hub_wizard(self.options)._build_report()
        sheet = self._workbook(report)
        self.assertEqual(sheet['A10'].value, '1')
        self.assertEqual(sheet['B10'].value, report['rows'][0]['name'].split('. ', 1)[-1])
        self.assertFalse(sheet['B10'].value.startswith('1. '))

    def test_precision_limit_rejects_without_saved_file(self):
        report = self.model._hub_wizard(self.options)._build_report()
        exports = self.env['baseer.tax.report.xlsx']
        before = exports.search_count([])
        report['rows'][0]['tax'] = Decimal('1000000000000.00')
        with self.assertRaises(UserError):
            exports._render_workbook(report)
        self.assertEqual(exports.search_count([]), before)
        report['rows'][0]['tax'] = Decimal('999999999999.99')
        sheet = self._workbook(report)
        self.assertEqual(Decimal(str(sheet['D10'].value)).quantize(Decimal('0.01')),
                         Decimal('999999999999.99'))

        report['currency'] = SimpleNamespace(name='TEST', decimal_places=3)
        report['rows'][0]['tax'] = Decimal('100000000000.000')
        with self.assertRaises(UserError):
            exports._render_workbook(report)
        report['rows'][0]['tax'] = Decimal('99999999999.999')
        sheet = self._workbook(report)
        self.assertEqual(Decimal(str(sheet['D10'].value)).quantize(Decimal('0.001')),
                         Decimal('99999999999.999'))

    def test_size_and_time_reject_before_exposing_download(self):
        report = self.model._hub_wizard(self.options)._build_report()
        exports = self.env['baseer.tax.report.xlsx']
        before = exports.search_count([])
        with patch.object(vat_xlsx_export, '_XLSX_LIMIT', 1):
            with self.assertRaises(UserError):
                exports._render_workbook(report)
        with self.assertRaises(UserError):
            exports._render_workbook(report, started_at=monotonic() - 4)
        self.assertEqual(exports.search_count([]), before)

    def test_one_hundred_journal_names_are_complete_on_a4_filters_sheet(self):
        report = self.model._hub_wizard(self.options)._build_report()
        report['wizard'] = SimpleNamespace(period_type='quarter', journal_ids=range(100))
        report['journal_names'] = ', '.join('=@Journal-%03d example' % index for index in range(100))
        payload = self.env['baseer.tax.report.xlsx']._render_workbook(report)
        workbook = load_workbook(BytesIO(payload))
        self.assertEqual(workbook.sheetnames, ['VAT Return', 'Filters'])
        self.assertIn('100 selected journals', workbook.active['B6'].value)
        self.assertIn('not the full VAT return', workbook.active['A7'].value)
        filters = workbook['Filters']
        reconstructed = ''.join(filters[f'B{row}'].value for row in range(6, filters.max_row + 1))
        self.assertEqual(reconstructed, report['journal_names'])
        self.assertEqual(str(filters.page_setup.paperSize), filters.PAPERSIZE_A4)
        self.assertEqual(filters.print_title_rows, '$1:$5')
        self.assertIn(f'$A$1:$B${filters.max_row}', filters.print_area)
        for row in filters:
            for cell in row:
                self.assertNotEqual(cell.data_type, 'f')
                self.assertIsNone(cell.hyperlink)

    def test_three_decimal_currency_matches_hub_pdf_and_xlsx(self):
        currency = self.env['res.currency'].create({
            'name': 'TVD', 'symbol': 'TVD', 'rounding': 0.001, 'active': True,
        })
        company = self.env['res.company'].create({
            'name': 'VAT three decimal test', 'country_id': self.env.ref('base.sa').id,
            'currency_id': currency.id,
        })
        company.account_fiscal_country_id = self.env.ref('base.sa')
        model = self.model.with_company(company).with_context(allowed_company_ids=[company.id])
        options = _options(company.id)
        values = {
            '-1(B)': Decimal('0.001'), '-1(T)': Decimal('0.001'),
            '7(B)': Decimal('0.004'), '7(T)': Decimal('0.004'),
        }
        exception = {
            'count': 1, 'amount': Decimal('-0.001'), 'domain': Domain.FALSE,
            'other_count': 1, 'other_amount': Decimal('0.001'),
            'other_domain': Domain.FALSE,
        }
        with patch.object(BaseerTaxReportWizard, '_tax_tag_expression',
                          side_effect=lambda expression, domain: values.get(
                              expression.formula, Decimal('0'))), \
                patch.object(BaseerTaxReportWizard, '_untagged_vat', return_value=exception):
            wizard = model._hub_wizard(options)
            report = wizard._build_report()
            hub = model.get_hub_report(options)
            report['visible_rows'] = report['rows']
            report['interactive'] = False
            pdf_table = self.env['ir.qweb']._render(
                'baseer_tax_report.tax_table', {'report_data': report},
            )
            payload = self.env['baseer.tax.report.xlsx']._render_workbook(report)

        self.assertEqual(currency.decimal_places, 3)
        self.assertEqual(report['rows'][0]['base'], Decimal('0.001'))
        self.assertEqual(report['rows'][0]['base_text'], '0.001')
        self.assertEqual(report['rows'][0]['tax_text'], '0.001')
        self.assertEqual(report['rows'][12]['tax_text'], '-0.003')
        self.assertEqual(report['rows'][12]['components']['tax'][1]['amount_text'], '-0.004')
        self.assertEqual(report['exception_amount_text'], '-0.001')
        self.assertEqual(report['other_tax_amount_text'], '0.001')
        self.assertEqual(hub['rows'][0]['base_text'], '0.001')
        self.assertEqual(hub['exception']['amount_text'], '-0.001')
        self.assertIn('0.001', pdf_table)
        self.assertIn('-0.003', pdf_table)
        self.assertIn('-0.001', pdf_table)
        sheet = load_workbook(BytesIO(payload)).active
        self.assertEqual(Decimal(str(sheet['C10'].value)), Decimal('0.001'))
        self.assertEqual(Decimal(str(sheet['D22'].value)), Decimal('-0.003'))
        self.assertEqual(Decimal(str(sheet['D28'].value)), Decimal('-0.001'))
        self.assertEqual(sheet['C10'].number_format, '#,##0.000')

    def test_model_download_rechecks_owner_company_and_accounting_group(self):
        group = self.env.ref('base.group_user') | self.env.ref('account.group_account_readonly')
        other_company = self.env['res.company'].create({
            'name': 'Other VAT XLSX company', 'country_id': self.env.ref('base.sa').id,
        })
        users = []
        for name in ('owner', 'other'):
            users.append(self.env['res.users'].create({
                'name': 'VAT XLSX %s' % name, 'login': 'vat.xlsx.%s' % name,
                'company_id': self.company.id,
                'company_ids': [Command.set((self.company | other_company).ids)],
                'group_ids': [Command.set(group.ids)],
            }))
        owner, other = users
        owner_model = self.model.with_user(owner).with_context(allowed_company_ids=[self.company.id])
        action = owner_model.export_hub_xlsx(self.options)
        self.assertEqual(action['target'], 'download')
        url = urlsplit(action['url'])
        record_id = int(url.path.rsplit('/', 1)[1])
        self.assertEqual(parse_qs(url.query)['company_id'], [str(self.company.id)])
        export = self.env['baseer.tax.report.xlsx'].browse(record_id)
        self.assertEqual(export.create_uid, owner)
        self.assertTrue(export.file_data)
        self.assertEqual(export.with_user(owner).with_context(
            allowed_company_ids=[self.company.id],
        )._download_record(record_id, self.company.id), export)
        with self.assertRaises(AccessError):
            export.with_user(other).with_context(
                allowed_company_ids=[self.company.id],
            )._download_record(record_id, self.company.id)
        with self.assertRaises(AccessError):
            export.with_user(owner).with_context(
                allowed_company_ids=[other_company.id],
            )._download_record(record_id, self.company.id)
        with self.assertRaises(AccessError):
            export.with_user(owner).with_context(
                allowed_company_ids=[self.company.id],
            )._download_record(record_id, other_company.id)
        plain_group = self.env.ref('base.group_user')
        owner.group_ids = [Command.set(plain_group.ids)]
        with self.assertRaises(AccessError):
            export.with_user(owner).with_context(
                allowed_company_ids=[self.company.id],
            )._download_record(record_id, self.company.id)


@tagged('post_install', '-at_install')
class TestVatXlsxHttp(HttpCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.company = cls.env.company
        cls.company.account_fiscal_country_id = cls.env.ref('base.sa')
        group = cls.env.ref('base.group_user') | cls.env.ref('account.group_account_readonly')
        cls.second_company = cls.env['res.company'].create({
            'name': 'VAT XLSX HTTP other company', 'country_id': cls.env.ref('base.sa').id,
        })
        cls.owner = cls.env['res.users'].create({
            'name': 'VAT XLSX HTTP owner', 'login': 'vat.xlsx.http.owner',
            'password': 'vat-xlsx-http-test', 'company_id': cls.company.id,
            'company_ids': [Command.set((cls.company | cls.second_company).ids)],
            'group_ids': [Command.set(group.ids)],
        })
        cls.other = cls.env['res.users'].create({
            'name': 'VAT XLSX HTTP other', 'login': 'vat.xlsx.http.other',
            'password': 'vat-xlsx-http-test', 'company_id': cls.company.id,
            'company_ids': [Command.set(cls.company.ids)],
            'group_ids': [Command.set(group.ids)],
        })

    def test_custom_and_generic_routes_enforce_owner_and_active_company(self):
        owner_model = self.env['baseer.tax.report.wizard'].with_user(self.owner).with_context(
            allowed_company_ids=[self.company.id],
        )
        action = owner_model.export_hub_xlsx(_options(self.company.id))
        record_id = int(urlsplit(action['url']).path.rsplit('/', 1)[1])
        generic = '/web/content?model=baseer.tax.report.xlsx&id=%s&field=file_data&download=true' % record_id
        self.authenticate(self.owner.login, 'vat-xlsx-http-test')
        owned = self.url_open(action['url'])
        self.assertEqual(owned.status_code, 200)
        self.assertIn('spreadsheetml.sheet', owned.headers['Content-Type'])
        self.assertIn('no-store', owned.headers['Cache-Control'])
        self.assertEqual(load_workbook(BytesIO(owned.content)).active['A10'].value, '1')
        self.assertEqual(self.url_open(action['url'] + '&field=other&model=res.users').content,
                         owned.content)
        self.assertEqual(self.url_open(action['url'].replace(
            'company_id=%s' % self.company.id,
            'company_id=%s' % self.second_company.id,
        )).status_code, 404)
        self.assertEqual(self.url_open('/baseer/tax/vat/export/2147483647?company_id=%s' %
                                       self.company.id).status_code, 404)
        self.assertEqual(self.url_open(generic).status_code, 404)
        self.authenticate(self.other.login, 'vat-xlsx-http-test')
        self.assertNotEqual(self.url_open(generic).status_code, 200)
        self.assertEqual(self.url_open(action['url']).status_code, 404)
        self.authenticate(self.owner.login, 'vat-xlsx-http-test')
        # Odoo's generic binary URL must obey the rule even after a company switch.
        self.owner.company_id = self.second_company
        self.assertNotEqual(self.url_open(generic).status_code, 200)
        self.assertEqual(self.url_open(action['url']).status_code, 404)
        self.opener.cookies.clear()
        self.assertNotEqual(self.url_open(action['url'], allow_redirects=False).status_code, 200)
