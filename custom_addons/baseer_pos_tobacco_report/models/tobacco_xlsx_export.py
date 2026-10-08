"""Private, bounded Excel rendering of the existing POS tobacco fee register."""

import logging
import re
from math import ceil
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from time import monotonic

from odoo import _, api, fields, models
from odoo.exceptions import AccessError, MissingError, UserError

from .tobacco_report import RIYADH
from .tobacco_report_wizard import ACCOUNT_GROUPS, POS_GROUPS


_logger = logging.getLogger(__name__)
_FILE_LIMIT = 5 * 1024 * 1024
_SECONDS_LIMIT = 10
_EXCEL_AMOUNT_LIMIT = Decimal('1000000000000')
_SAR_QUANTUM = Decimal('0.01')
_ILLEGAL_XML = re.compile(r'[\x00-\x08\x0b-\x0c\x0e-\x1f]')


class BaseerPosTobaccoXlsx(models.TransientModel):
    _name = 'baseer.pos.tobacco.xlsx'
    _description = 'Private tobacco fee XLSX download'
    _transient_max_hours = 1

    company_id = fields.Many2one('res.company', required=True, readonly=True)
    period = fields.Selection([('month', 'Month'), ('custom', 'Custom')], required=True, readonly=True)
    month = fields.Integer(required=True, readonly=True)
    year = fields.Integer(required=True, readonly=True)
    date_from = fields.Date(required=True, readonly=True)
    date_to = fields.Date(required=True, readonly=True)
    show_products = fields.Boolean(readonly=True)
    # Text deliberately prevents /web/content from serving another user's file.
    file_data = fields.Text(readonly=True)
    file_name = fields.Char(readonly=True)

    def check_access(self, operation):
        result = super().check_access(operation)
        user = self.env.user
        if not any(user.has_group(group) for group in ACCOUNT_GROUPS) or not any(
            user.has_group(group) for group in POS_GROUPS
        ):
            raise AccessError(_('The tobacco fee export requires both Accounting and POS read access.'))
        return result

    @api.model
    def _download_record(self, export_id, company_id):
        if type(export_id) is not int or type(company_id) is not int:
            raise AccessError(_('Invalid tobacco fee export reference.'))
        if company_id not in self.env.companies.ids:
            raise AccessError(_('The export company is not active.'))
        export = self.browse(export_id).exists()
        if not export:
            raise MissingError(_('The tobacco fee export has expired. Generate it again.'))
        export.check_access('read')
        if export.create_uid.id != self.env.uid or export.company_id.id != company_id:
            raise AccessError(_('The tobacco fee export is not available.'))
        company, values = self.env['baseer.pos.tobacco.report.wizard']._hub_wizard_values({
            'company_id': company_id, 'period': export.period,
            'month': export.month, 'year': export.year,
            'date_from': fields.Date.to_string(export.date_from),
            'date_to': fields.Date.to_string(export.date_to),
            'show_products': export.show_products, 'page': 1,
        })
        wizard = self.env['baseer.pos.tobacco.report.wizard'].with_company(company).new(values)
        wizard.check_access('read')
        wizard._check_report_access()
        wizard._check_period()
        if not export.file_data or not export.file_name:
            raise MissingError(_('The tobacco fee export has expired. Generate it again.'))
        return export

    @api.model
    def _render_workbook(self, report, started_at=None):
        """Serialize source Decimal rows, not screen strings or recalculated fees."""
        started_at = started_at if started_at is not None else monotonic()
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
            from openpyxl.worksheet.page import PageMargins
        except ImportError as error:
            _logger.warning('Tobacco XLSX rejected: openpyxl unavailable')
            raise UserError(_('Excel export is unavailable on this server.')) from error

        source = report.get('_export')
        if not source or len(source['rows']) != len(report['rows']) or len(source['exceptions']) != len(report['exceptions']):
            raise UserError(_('The tobacco fee export requires complete source data.'))
        is_rtl = report['is_rtl']
        show_products = report['show_products']
        columns = 7 if show_products else 6
        last_col = 'G' if show_products else 'F'
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'رسوم التبغ' if is_rtl else 'Tobacco Fees'
        sheet.sheet_view.rightToLeft = is_rtl
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = 'D11' if show_products else 'C11'
        widths = [20, 22, 14, 42, 15, 15, 17] if show_products else [20, 22, 14, 15, 15, 17]
        for index, width in enumerate(widths, 1):
            sheet.column_dimensions[chr(64 + index)].width = width

        ink, muted, negative = '243248', 'AAB7C4', 'D3262E'
        header_fill = PatternFill('solid', fgColor='EDF1F4')
        total_fill = PatternFill('solid', fgColor='DCE0E3')
        warning_fill = PatternFill('solid', fgColor='FFF4DF')
        separator = Side(style='hair', color='DCE3E9')

        def literal(target_sheet, address, value, *, bold=False, fill=None, wrap=False):
            cell = target_sheet[address]
            cell.value = _ILLEGAL_XML.sub('', str(value if value is not None else ''))[:32767]
            cell.data_type = 's'
            cell.hyperlink = None
            cell.font = Font(name='Arial', size=10, bold=bold, color=ink)
            cell.alignment = Alignment(horizontal='right' if is_rtl else 'left',
                                       vertical='center', wrap_text=wrap)
            if fill:
                cell.fill = fill
            return cell

        def money(target_sheet, address, value, *, bold=False, fill=None, outflow=False):
            if type(value) is not Decimal or not value.is_finite() or abs(value) >= _EXCEL_AMOUNT_LIMIT:
                _logger.warning('Tobacco XLSX rejected: unsafe numeric precision')
                raise UserError(_('A tobacco fee amount exceeds safe Excel precision. Split or correct the source.'))
            try:
                rounded = value.quantize(_SAR_QUANTUM)
                preserved = Decimal(str(float(value))).quantize(_SAR_QUANTUM)
            except (ValueError, ArithmeticError, OverflowError) as error:
                raise UserError(_('A tobacco fee amount cannot be exported safely to Excel.')) from error
            if rounded != value or preserved != value:
                raise UserError(_('A tobacco fee amount cannot be exported safely to Excel.'))
            cell = target_sheet[address]
            cell.value = value
            cell.number_format = '#,##0.00'
            cell.font = Font(name='Arial', size=10, bold=bold,
                             color=muted if value == 0 else negative if value < 0 or outflow else ink)
            cell.alignment = Alignment(horizontal='right', vertical='center')
            if fill:
                cell.fill = fill
            return cell

        def typed_date(target_sheet, address, value):
            if type(value) is not datetime or value.tzinfo is not None:
                raise UserError(_('Invalid tobacco fee event date.'))
            cell = target_sheet[address]
            cell.value = value
            cell.number_format = 'yyyy-mm-dd hh:mm'
            cell.font = Font(name='Arial', size=10, color=ink)
            cell.alignment = Alignment(horizontal='center', vertical='center')
            return cell

        def wrapped_height(*items):
            """Keep long POS references and products inside their printed rows."""
            line_count = max(
                sum(max(1, ceil(len(part) / max(1, width - 5)))
                    for part in str(value).splitlines() or [''])
                for value, width in items
            )
            height = max(25, 10 + 16 * line_count)
            if height > 409:
                raise UserError(_('A POS label is too long to print safely in Excel. Correct the source.'))
            return height

        title = 'كشف رسوم التبغ من نقاط البيع' if is_rtl else 'POS Tobacco Fee Register'
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
        literal(sheet, 'A1', title, bold=True).font = Font(name='Arial', size=15, bold=True, color=ink)
        sheet.row_dimensions[1].height = 31
        labels = {
            'company': 'الشركة' if is_rtl else 'Company',
            'period': 'الفترة' if is_rtl else 'Period',
            'currency': 'العملة' if is_rtl else 'Currency',
            'generated': 'تاريخ التوليد (الرياض)' if is_rtl else 'Generated (Riyadh)',
            'products': 'إظهار المنتج' if is_rtl else 'Products',
            'yes': 'نعم' if is_rtl else 'Yes',
            'no': 'لا' if is_rtl else 'No',
            'source': 'المصدر' if is_rtl else 'Source',
        }
        literal(sheet, 'A2', labels['company'], bold=True)
        sheet.merge_cells(start_row=2, start_column=2, end_row=2, end_column=columns)
        literal(sheet, 'B2', report['company'].name)
        literal(sheet, 'A3', labels['period'], bold=True)
        sheet['B3'] = report['wizard'].date_from
        sheet['B3'].number_format = 'yyyy-mm-dd'
        sheet['B3'].alignment = Alignment(horizontal='right' if is_rtl else 'left',
                                          vertical='center')
        literal(sheet, 'C3', 'إلى' if is_rtl else 'To', bold=True)
        sheet['D3'] = report['wizard'].date_to
        sheet['D3'].number_format = 'yyyy-mm-dd'
        sheet['D3'].alignment = Alignment(horizontal='right' if is_rtl else 'left',
                                          vertical='center')
        literal(sheet, 'A4', labels['currency'], bold=True)
        literal(sheet, 'B4', report['currency'])
        literal(sheet, 'C4', labels['products'], bold=True)
        literal(sheet, 'D4', labels['yes'] if show_products else labels['no'])
        literal(sheet, 'A5', labels['generated'], bold=True)
        sheet['B5'] = datetime.now(RIYADH).replace(tzinfo=None)
        sheet['B5'].number_format = 'yyyy-mm-dd hh:mm'
        sheet['B5'].alignment = Alignment(horizontal='right' if is_rtl else 'left',
                                          vertical='center')
        literal(sheet, 'A6', labels['source'], bold=True)
        sheet.merge_cells(start_row=6, start_column=2, end_row=6, end_column=columns)
        literal(sheet, 'B6', 'أوامر POS المكتملة؛ الرسوم معاد احتسابها من المصدر' if is_rtl
                else 'Completed POS orders; fees recalculated from source')
        sheet.merge_cells(start_row=7, start_column=1, end_row=7, end_column=columns)
        literal(sheet, 'A7', 'الفروق عن الإجماليات المحفوظة في ورقة منفصلة' if is_rtl
                else 'Differences from saved order totals are shown on a separate sheet',
                fill=warning_fill)
        sheet.row_dimensions[7].height = 24
        sheet.row_dimensions[9].height = 8
        headers = ([('التاريخ والوقت' if is_rtl else 'Date and Time'),
                    ('طلب نقطة البيع' if is_rtl else 'POS Order'),
                    ('النوع' if is_rtl else 'Type')]
                   + ([('المنتج والكمية' if is_rtl else 'Product and Quantity')] if show_products else [])
                   + [('خارج' if is_rtl else 'Out'), ('داخل' if is_rtl else 'In'),
                      ('الرصيد التراكمي' if is_rtl else 'Running Total')])
        for index, value in enumerate(headers, 1):
            literal(sheet, f'{chr(64 + index)}10', value, bold=True, fill=header_fill)
        sheet.row_dimensions[10].height = 26
        sheet.merge_cells(start_row=11, start_column=1, end_row=11, end_column=columns - 1)
        literal(sheet, 'A11', 'رصيد بداية الفترة' if is_rtl else 'Opening Balance', bold=True)
        money(sheet, f'{last_col}11', source['opening'], bold=True)

        for number, row in enumerate(source['rows'], 12):
            typed_date(sheet, f'A{number}', row['date'])
            literal(sheet, f'B{number}', row['order'], wrap=True)
            literal(sheet, f'C{number}', row['type'])
            if show_products:
                literal(sheet, f'D{number}', row['products'], wrap=True)
            for letter, key in zip(('E', 'F', 'G') if show_products else ('D', 'E', 'F'),
                                   ('debit', 'credit', 'running')):
                money(sheet, f'{letter}{number}', row[key], outflow=key == 'debit')
            for cell in sheet[number]:
                cell.border = Border(bottom=separator)
            wrapped_items = [(row['order'], widths[1])]
            if show_products:
                wrapped_items.append((row['products'], widths[3]))
            sheet.row_dimensions[number].height = wrapped_height(*wrapped_items)
        total_row = 12 + len(source['rows'])
        sheet.merge_cells(start_row=total_row, start_column=1, end_row=total_row,
                          end_column=columns - 3)
        literal(sheet, f'A{total_row}', 'مجموع الحركة' if is_rtl else 'Period Movement',
                bold=True, fill=total_fill)
        for letter, key in zip(('E', 'F', 'G') if show_products else ('D', 'E', 'F'),
                               ('debit_total', 'credit_total', 'closing')):
            money(sheet, f'{letter}{total_row}', source[key], bold=True, fill=total_fill,
                  outflow=key == 'debit_total')
        sheet.row_dimensions[total_row].height = 26
        closing_row = total_row + 1
        sheet.merge_cells(start_row=closing_row, start_column=1, end_row=closing_row,
                          end_column=columns - 1)
        literal(sheet, f'A{closing_row}', 'الرصيد الختامي' if is_rtl else 'Closing Balance',
                bold=True)
        money(sheet, f'{last_col}{closing_row}', source['closing'], bold=True)
        sheet.row_dimensions[closing_row].height = 26

        exceptions = workbook.create_sheet('الاستثناءات' if is_rtl else 'Exceptions')
        exceptions.sheet_view.rightToLeft = is_rtl
        exceptions.sheet_view.showGridLines = False
        for letter, width in {'A': 23, 'B': 34, 'C': 26, 'D': 26}.items():
            exceptions.column_dimensions[letter].width = width
        exceptions.merge_cells('A1:D1')
        literal(exceptions, 'A1', 'فروق إجمالي طلبات POS' if is_rtl else 'POS Order Total Differences',
                bold=True).font = Font(name='Arial', size=15, bold=True, color=ink)
        exceptions.row_dimensions[1].height = 31
        exceptions.merge_cells('A2:D2')
        literal(exceptions, 'A2', report['company'].name)
        exceptions['A3'] = report['wizard'].date_from
        exceptions['A3'].number_format = 'yyyy-mm-dd'
        exceptions['B3'] = report['wizard'].date_to
        exceptions['B3'].number_format = 'yyyy-mm-dd'
        exceptions.merge_cells('A4:D4')
        literal(exceptions, 'A4', 'هذه الفروق معلومات مراجعة؛ لا تغيّر رسوم الكشف' if is_rtl
                else 'Review differences only; register fees are not changed', fill=warning_fill)
        for index, value in enumerate((
            'التاريخ والوقت' if is_rtl else 'Date and Time',
            'طلب نقطة البيع' if is_rtl else 'POS Order',
            'الإجمالي المعاد احتسابه' if is_rtl else 'Recalculated Total',
            'الإجمالي المحفوظ' if is_rtl else 'Saved Total',
        ), 1):
            literal(exceptions, f'{chr(64 + index)}6', value, bold=True, fill=header_fill)
        exceptions.row_dimensions[6].height = 26
        for number, row in enumerate(source['exceptions'], 7):
            typed_date(exceptions, f'A{number}', row['date'])
            literal(exceptions, f'B{number}', row['order'])
            money(exceptions, f'C{number}', row['computed_total'])
            money(exceptions, f'D{number}', row['saved_total'])
            for cell in exceptions[number]:
                cell.border = Border(bottom=separator)
            exceptions.row_dimensions[number].height = 23
        if not source['exceptions']:
            exceptions.merge_cells('A7:D7')
            literal(exceptions, 'A7', 'لا توجد فروق' if is_rtl else 'No differences')

        def configure_print(target_sheet, area, title_rows, orientation):
            target_sheet.print_title_rows = title_rows
            target_sheet.print_area = area
            target_sheet.print_options.horizontalCentered = True
            target_sheet.sheet_properties.pageSetUpPr.fitToPage = True
            target_sheet.page_setup.orientation = orientation
            target_sheet.page_setup.paperSize = target_sheet.PAPERSIZE_A4
            target_sheet.page_setup.fitToWidth = 1
            target_sheet.page_setup.fitToHeight = 0
            target_sheet.page_margins = PageMargins(left=0.3, right=0.3, top=0.45,
                                                    bottom=0.45, header=0.2, footer=0.2)
            target_sheet.oddFooter.center.text = 'Page &P / &N'

        configure_print(sheet, f'A1:{last_col}{closing_row}', '1:10',
                        'landscape' if show_products else 'portrait')
        configure_print(exceptions, f'A1:D{max(7, 6 + len(source["exceptions"]))}', '1:6',
                        'portrait')
        output = BytesIO()
        workbook.save(output)
        payload = output.getvalue()
        duration = monotonic() - started_at
        if len(payload) > _FILE_LIMIT or duration > _SECONDS_LIMIT:
            _logger.warning('Tobacco XLSX rejected: bytes=%s elapsed_ms=%s',
                            len(payload), int(duration * 1000))
            raise UserError(_('Tobacco Excel export exceeded its safe size or time limit. Split the period.'))
        _logger.info('Tobacco XLSX prepared: rows=%s exceptions=%s bytes=%s elapsed_ms=%s',
                     len(source['rows']), len(source['exceptions']), len(payload), int(duration * 1000))
        return payload
