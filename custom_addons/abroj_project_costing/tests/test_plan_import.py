import base64
from io import BytesIO

from openpyxl import Workbook, load_workbook

from odoo.exceptions import UserError, ValidationError
from odoo.tests.common import TransactionCase, tagged


@tagged('post_install', '-at_install')
class TestAbrojCostPlanImport(TransactionCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.company = cls.env['res.company'].create({
            'name': 'Plan import test company',
            'abroj_project_costing_enabled': True,
        })
        cls.category = cls.env['abroj.cost.category'].with_company(cls.company).create({
            'name': 'التشطيبات',
            'company_id': cls.company.id,
        })

    def _project(self):
        return self.env['abroj.cost.project'].with_company(self.company).create({
            'name': 'Plan import test project',
            'company_id': self.company.id,
            'agreement_amount': 10000,
        })

    def _wizard(self, project, rows):
        content = '\n'.join([
            'category,name,pricing_method,quantity,material_unit_cost,auxiliary_unit_cost,labor_unit_cost,inclusive_unit_cost,lump_sum_cost,notes',
            *rows,
        ]).encode()
        return self.env['abroj.cost.plan.import.wizard'].with_company(self.company).create({
            'project_id': project.id,
            'company_id': self.company.id,
            'import_file': base64.b64encode(content),
            'import_filename': 'plan.csv',
        })

    def test_csv_import_creates_only_plan_lines(self):
        project = self._project()
        material_count_before = self.env['abroj.cost.material'].search_count([
            ('company_id', '=', self.company.id),
        ])
        wizard = self._wizard(project, [
            'التشطيبات,بند مقطوعية,lump_sum,1,0,0,0,0,1500,ملاحظة',
            'التشطيبات,بند تفصيلي,detailed,2,100,25,50,0,0,',
        ])
        action = wizard.action_import_plan()
        lines = project.plan_line_ids.sorted('sequence')
        self.assertEqual(len(lines), 2)
        self.assertEqual(lines[0].pricing_method, 'lump_sum')
        self.assertEqual(lines[0].estimated_total, 1500)
        self.assertEqual(lines[1].estimated_total, 350)
        self.assertEqual(action['tag'], 'display_notification')
        self.assertEqual(self.env['abroj.cost.material'].search_count([
            ('company_id', '=', self.company.id),
        ]), material_count_before)

    def test_xlsx_import_creates_plan_lines(self):
        project = self._project()
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'بنود الدراسة'
        sheet.append([
            'category', 'name', 'pricing_method', 'quantity',
            'material_unit_cost', 'auxiliary_unit_cost', 'labor_unit_cost',
            'inclusive_unit_cost', 'lump_sum_cost', 'notes',
        ])
        sheet.append(['التشطيبات', 'بند Excel', 'supply_install', 2, 0, 0, 0, 375, 0, ''])
        content = BytesIO()
        workbook.save(content)
        wizard = self.env['abroj.cost.plan.import.wizard'].with_company(self.company).create({
            'project_id': project.id,
            'company_id': self.company.id,
            'import_file': base64.b64encode(content.getvalue()),
            'import_filename': 'plan.xlsx',
        })
        wizard.action_import_plan()
        self.assertEqual(project.plan_line_ids.estimated_total, 750)

    def test_invalid_row_does_not_create_partial_plan(self):
        project = self._project()
        wizard = self._wizard(project, [
            'التشطيبات,بند صحيح,lump_sum,1,0,0,0,0,1500,',
            'فئة غير موجودة,بند مرفوض,lump_sum,1,0,0,0,0,900,',
        ])
        with self.assertRaises(UserError):
            wizard.action_import_plan()
        self.assertFalse(project.plan_line_ids)

    def test_import_rejects_existing_plan_to_prevent_duplicates(self):
        project = self._project()
        self.env['abroj.cost.plan.line'].with_company(self.company).create({
            'project_id': project.id,
            'category_id': self.category.id,
            'name': 'بند موجود',
            'pricing_method': 'lump_sum',
            'lump_sum_cost': 10,
        })
        wizard = self._wizard(project, [
            'التشطيبات,بند جديد,lump_sum,1,0,0,0,0,1500,',
        ])
        with self.assertRaises(UserError):
            wizard.action_import_plan()
        self.assertEqual(len(project.plan_line_ids), 1)

    def test_import_rechecks_plan_after_validation(self):
        """A later import must see a plan created after its file was read."""
        project = self._project()
        wizard = self._wizard(project, [
            'التشطيبات,بند متأخر,lump_sum,1,0,0,0,0,1500,',
        ])
        wizard._validated_values()
        self.env['abroj.cost.plan.line'].with_company(self.company).create({
            'project_id': project.id,
            'category_id': self.category.id,
            'name': 'بند استيراد آخر',
            'pricing_method': 'lump_sum',
            'lump_sum_cost': 10,
        })
        with self.assertRaises(UserError):
            wizard._lock_empty_project_plan()
        self.assertEqual(len(project.plan_line_ids), 1)

    def test_csv_rejects_row_501_before_import(self):
        project = self._project()
        wizard = self._wizard(project, [
            'التشطيبات,بند %s,lump_sum,1,0,0,0,0,1,' % index
            for index in range(1, 502)
        ])
        with self.assertRaises(ValidationError):
            wizard.action_import_plan()
        self.assertFalse(project.plan_line_ids)

    def test_export_returns_download_action(self):
        project = self._project()
        self.env['abroj.cost.plan.line'].with_company(self.company).create({
            'project_id': project.id,
            'category_id': self.category.id,
            'name': 'بند للتصدير',
            'pricing_method': 'lump_sum',
            'lump_sum_cost': 1500,
        })
        action = project.action_export_plan()
        self.assertEqual(action['type'], 'ir.actions.act_url')
        self.assertIn('download=true', action['url'])

    def test_export_escapes_spreadsheet_formula_text(self):
        project = self._project()
        self.env['abroj.cost.plan.line'].with_company(self.company).create({
            'project_id': project.id,
            'category_id': self.category.id,
            'name': '=2+2',
            'pricing_method': 'lump_sum',
            'lump_sum_cost': 1500,
            'notes': '@unsafe',
        })
        wizard = self.env['abroj.cost.plan.import.wizard'].with_company(self.company).create({
            'project_id': project.id,
            'company_id': self.company.id,
        })
        wizard.action_export_plan()
        workbook = load_workbook(BytesIO(base64.b64decode(wizard.export_file)), data_only=False)
        sheet = workbook['بنود الدراسة']
        self.assertEqual(sheet['B2'].data_type, 's')
        self.assertEqual(sheet['B2'].value, "'=2+2")
        self.assertEqual(sheet['J2'].value, "'@unsafe")

    def test_tree_import_accepts_sections_then_priced_leaves(self):
        project = self._project()
        content = '\n'.join([
            'category,name,pricing_method,quantity,material_unit_cost,auxiliary_unit_cost,labor_unit_cost,inclusive_unit_cost,lump_sum_cost,notes,path,node_kind',
            ',المطبخ,detailed,0,0,0,0,0,0,,المطبخ,section',
            ',سباكة,detailed,0,0,0,0,0,0,,المطبخ/سباكة,section',
            'التشطيبات,تمديدات,lump_sum,1,0,0,0,0,200,,المطبخ/سباكة/تمديدات,item',
        ]).encode()
        wizard = self.env['abroj.cost.plan.import.wizard'].with_company(self.company).create({
            'project_id': project.id,
            'company_id': self.company.id,
            'import_file': base64.b64encode(content),
            'import_filename': 'tree.csv',
        })
        wizard.action_import_plan()
        kitchen = project.plan_line_ids.filtered(lambda line: line.name == 'المطبخ')
        plumbing = project.plan_line_ids.filtered(lambda line: line.name == 'سباكة')
        leaf = project.plan_line_ids.filtered(lambda line: line.name == 'تمديدات')
        self.assertEqual(kitchen.node_kind, 'section')
        self.assertEqual(plumbing.parent_id, kitchen)
        self.assertEqual(leaf.parent_id, plumbing)
        self.assertEqual(project.estimated_total, 200)

    def test_download_template_is_importable(self):
        project = self._project()
        template_wizard = self.env['abroj.cost.plan.import.wizard'].with_company(self.company).create({
            'project_id': project.id,
            'company_id': self.company.id,
        })
        import_wizard = self.env['abroj.cost.plan.import.wizard'].with_company(self.company).create({
            'project_id': project.id,
            'company_id': self.company.id,
            'import_file': base64.b64encode(template_wizard._template_workbook()),
            'import_filename': 'template.xlsx',
        })
        import_wizard.action_import_plan()
        kitchen = project.plan_line_ids.filtered(lambda line: line.name == 'المطبخ')
        example = project.plan_line_ids.filtered(lambda line: line.name == 'مثال بند')
        self.assertEqual(kitchen.node_kind, 'section')
        self.assertEqual(example.parent_id, kitchen)

    def test_template_download_uses_an_existing_company_category(self):
        self.category.write({'name': 'كهرباء الشركة', 'sequence': 0})
        project = self._project()
        wizard = self.env['abroj.cost.plan.import.wizard'].with_company(self.company).create({
            'project_id': project.id, 'company_id': self.company.id,
        })
        action = wizard.action_download_template()
        self.assertEqual(action['type'], 'ir.actions.act_url')
        self.assertEqual(wizard.export_filename, 'abroj_cost_plan_template.xlsx')
        workbook = load_workbook(BytesIO(base64.b64decode(wizard.export_file)))
        self.assertEqual(workbook['بنود الدراسة']['A3'].value, self.category.name)
        self.assertEqual(workbook['بنود الدراسة']['L2'].value, 'section')
        self.assertEqual(workbook['بنود الدراسة']['L3'].value, 'item')
        self.assertGreater(workbook['التعليمات'].max_row, 8)
        wizard.write({'import_file': wizard.export_file, 'import_filename': wizard.export_filename})
        wizard.action_import_plan()
        self.assertEqual(project.estimated_total, 1500)

    def test_template_without_active_categories_contains_only_a_section(self):
        project = self._project()
        self.env['abroj.cost.category'].search([('company_id', '=', self.company.id)]).write({'active': False})
        wizard = self.env['abroj.cost.plan.import.wizard'].with_company(self.company).create({
            'project_id': project.id, 'company_id': self.company.id,
        })
        wizard.action_download_template()
        wizard.write({'import_file': wizard.export_file, 'import_filename': wizard.export_filename})
        wizard.action_import_plan()
        self.assertEqual(len(project.plan_line_ids), 1)
        self.assertEqual(project.plan_line_ids.node_kind, 'section')
        self.assertEqual(project.estimated_total, 0)

    def test_template_keeps_formula_like_category_name_literal_and_importable(self):
        self.category.write({'name': '=كهرباء', 'sequence': 0})
        project = self._project()
        wizard = self.env['abroj.cost.plan.import.wizard'].with_company(self.company).create({
            'project_id': project.id, 'company_id': self.company.id,
        })
        wizard.action_download_template()
        workbook = load_workbook(BytesIO(base64.b64decode(wizard.export_file)), data_only=False)
        self.assertEqual(workbook['بنود الدراسة']['A3'].data_type, 's')
        self.assertEqual(workbook['بنود الدراسة']['A3'].value, self.category.name)
        wizard.write({'import_file': wizard.export_file, 'import_filename': wizard.export_filename})
        wizard.action_import_plan()
        self.assertEqual(project.plan_line_ids.filtered(lambda line: line.node_kind == 'item').category_id, self.category)

    def test_export_reimport_preserves_dragged_tree_order_and_amounts(self):
        source = self._project()
        Line = self.env['abroj.cost.plan.line'].with_company(self.company)
        kitchen = Line.create({'project_id': source.id, 'name': 'المطبخ', 'node_kind': 'section'})
        lounge = Line.create({'project_id': source.id, 'name': 'الصالة', 'node_kind': 'section'})
        plumbing = Line.create({'project_id': source.id, 'parent_id': kitchen.id, 'name': 'سباكة', 'node_kind': 'section'})
        electricity = Line.create({'project_id': source.id, 'parent_id': kitchen.id, 'name': 'كهرباء', 'node_kind': 'section'})
        for parent, name, amount in [(plumbing, 'تمديدات', 200), (electricity, 'تمديدات', 350), (lounge, 'دهان', 450)]:
            Line.create({
                'project_id': source.id, 'parent_id': parent.id, 'name': name,
                'category_id': self.category.id, 'pricing_method': 'lump_sum', 'lump_sum_cost': amount,
            })
        source.action_reorder_plan_section(kitchen.id, lounge.id)
        source.action_reorder_plan_section(plumbing.id, electricity.id)
        expected_names = ['الصالة', 'دهان', 'المطبخ', 'كهرباء', 'تمديدات', 'سباكة', 'تمديدات']
        self.assertEqual([row['name'] for row in source.get_study_tree_data()], expected_names)
        wizard = self.env['abroj.cost.plan.import.wizard'].with_company(self.company).create({
            'project_id': source.id, 'company_id': self.company.id,
        })
        wizard.action_export_plan()
        workbook = load_workbook(BytesIO(base64.b64decode(wizard.export_file)))
        self.assertEqual([row[1] for row in list(workbook['بنود الدراسة'].values)[1:]], expected_names)
        target = self._project()
        importer = self.env['abroj.cost.plan.import.wizard'].with_company(self.company).create({
            'project_id': target.id, 'company_id': self.company.id,
            'import_file': wizard.export_file, 'import_filename': wizard.export_filename,
        })
        importer.action_import_plan()
        self.assertEqual([row['name'] for row in target.get_study_tree_data()], expected_names)
        self.assertEqual(target.estimated_total, source.estimated_total)
        self.assertEqual(target.estimated_total, 1000)
        self.assertEqual(
            {wizard._plan_path(line): (line.node_kind, line.estimated_total) for line in source.plan_line_ids},
            {wizard._plan_path(line): (line.node_kind, line.estimated_total) for line in target.plan_line_ids},
        )
