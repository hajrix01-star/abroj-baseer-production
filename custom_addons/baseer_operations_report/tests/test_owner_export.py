"""Focused contract tests for the private unified owner-report export."""

import json
from decimal import Decimal
from io import BytesIO
import os
from pathlib import Path
import re
from unittest.mock import patch

from openpyxl import load_workbook

from odoo import Command
from odoo.exceptions import AccessError, ValidationError
from odoo.tests.common import HttpCase, TransactionCase, tagged

from ..report.owner_report import _owner_print_rows
from . import test_owner_cash


@tagged('post_install', '-at_install', 'reports_unified_export')
class TestOwnerUnifiedExport(TransactionCase):
    """The export is a typed projection of the already-authoritative snapshot."""

    _statement = test_owner_cash.TestOwnerCash._statement

    def setUp(self):
        super().setUp()
        self.env['res.lang']._activate_lang('ar_001')
        self.company = self.env.company
        self.report = self.env['baseer.operations.report']
        self.exports = self.env['baseer.report.export']
        self.expense = self.env['account.account'].with_company(self.company).create({
            'code': '959801', 'name': 'Unified export operating expense',
            'account_type': 'expense', 'company_ids': [Command.set(self.company.ids)],
        })
        self.bank = self.env['account.journal'].with_company(self.company).search([
            ('company_id', '=', self.company.id), ('type', '=', 'bank'),
        ], limit=1)
        self.assertTrue(self.bank and self.bank.default_account_id)
        self.partner = self.env['res.partner'].with_company(self.company).create({
            'name': 'Unified export statement partner',
        })
        self.filters = {
            'company_id': self.company.id, 'date_from': '2026-06-01',
            'date_to': '2026-06-30', 'journal_ids': [],
        }

    @staticmethod
    def _evidence(name, payload):
        """The isolated visual runner may request artifacts; ordinary tests write none."""
        target = os.environ.get('BASEER_EXPORT_EVIDENCE_DIR')
        if target and os.path.isdir(target):
            Path(target, name).write_bytes(payload)

    def test_actual_document_preserves_authoritative_decimal_cells_and_downloads(self):
        self._statement('40.00')
        snapshot = self.report.get_owner_snapshot(self.filters)
        self.assertTrue(snapshot['complete'])
        document = self.exports._document('operations_gross', self.filters)
        source_rows = _owner_print_rows(snapshot['periods'])
        self.assertEqual(len(document['rows']), len(source_rows))
        for exported, source in zip(document['rows'], source_rows):
            self.assertEqual(exported['cells']['label'], source['label'])
            for period_key, cell in source['amounts'].items():
                self.assertEqual(exported['cells'][period_key],
                                 Decimal(cell['amount'].replace(',', '')))
                if cell.get('percentage') is not None:
                    self.assertEqual(exported['cells'][period_key + '_percentage'],
                                     Decimal(cell['percentage'][:-1]).scaleb(-2))

        payload, mimetype, name = self.exports._export('operations_gross', self.filters, 'xlsx')
        self.assertIn('spreadsheetml.sheet', mimetype)
        self.assertTrue(name.endswith('.xlsx'))
        sheet = load_workbook(BytesIO(payload)).active
        self.assertEqual(sheet.sheet_view.rightToLeft, False)
        self.assertEqual(str(sheet.page_setup.paperSize), str(sheet.PAPERSIZE_A4))
        self.assertTrue(any(cell.value == Decimal('40.00')
                            for row in sheet.iter_rows() for cell in row))
        self._evidence('owner-export.xlsx', payload)

        payload, mimetype, name = self.exports._export('operations_gross', self.filters, 'pdf')
        self.assertEqual(mimetype, 'application/pdf')
        self.assertTrue(name.endswith('.pdf'))
        self.assertTrue(payload.startswith(b'%PDF-'))
        try:
            from pypdf import PdfReader
        except ImportError:
            from PyPDF2 import PdfReader
        pages = PdfReader(BytesIO(payload)).pages
        self.assertTrue(pages)
        self.assertAlmostEqual(float(pages[0].mediabox.width), 595, delta=4)
        self.assertAlmostEqual(float(pages[0].mediabox.height), 842, delta=4)
        self.assertRegex('\n'.join(page.extract_text() or '' for page in pages), r'\b1\s*/\s*1\b')
        # Body and separately-rendered page footer must use the packaged font.
        for page in pages:
            for reference in page['/Resources']['/Font'].values():
                font = reference.get_object()
                self.assertIn('IBMPlexSansArabic', str(font['/BaseFont']))
                descendant = font.get('/DescendantFonts', [reference])[0].get_object()
                descriptor = descendant['/FontDescriptor'].get_object()
                self.assertTrue(any(key in descriptor for key in ('/FontFile', '/FontFile2', '/FontFile3')))
        self._evidence('owner-export.pdf', payload)

    def test_arabic_and_comparison_files_render_with_safe_failure(self):
        self._statement('40.00')
        engine = self.exports.with_context(lang='ar_001')
        data, _mime, _name = engine._export('operations_gross', self.filters, 'xlsx')
        sheet = load_workbook(BytesIO(data)).active
        self.assertTrue(sheet.sheet_view.rightToLeft)
        self.assertTrue(all(cell.hyperlink is None for row in sheet for cell in row))
        self._evidence('owner-export-ar.xlsx', data)
        data, _mime, _name = engine._export('operations_gross', self.filters, 'pdf')
        self.assertTrue(data.startswith(b'%PDF-'))
        self._evidence('owner-export-ar.pdf', data)
        comparison = {**self.filters, 'comparison': {'kind': 'previous_period', 'count': 3, 'order': 'descending'}}
        data, _mime, _name = engine._export('operations_gross', comparison, 'pdf')
        self._evidence('owner-export-comparison-ar.pdf', data)
        try:
            from pypdf import PdfReader
        except ImportError:
            from PyPDF2 import PdfReader
        pages = PdfReader(BytesIO(data)).pages
        self.assertAlmostEqual(float(pages[0].mediabox.width), 842, delta=4)
        self.assertAlmostEqual(float(pages[0].mediabox.height), 595, delta=4)
        for page in pages:
            self.assertFalse(any(annotation.get_object().get('/Subtype') == '/Link'
                                 for annotation in page.get('/Annots', [])))
        with patch.object(type(self.env['ir.actions.report']), '_run_wkhtmltopdf', side_effect=OSError('private temp path')):
            with self.assertRaises(ValidationError) as caught:
                engine._export('operations_gross', self.filters, 'pdf')
        self.assertNotIn('private temp path', str(caught.exception))

    def test_rejects_incomplete_foreign_and_tampered_filters(self):
        with patch.object(type(self.report), 'get_owner_snapshot', return_value={
                'complete': False, 'company_id': self.company.id}):
            with self.assertRaises(AccessError):
                self.exports._document('operations_gross', self.filters)
        other = self.env['res.company'].create({'name': 'Unified export foreign company'})
        with self.assertRaises(AccessError):
            self.exports._document('operations_gross', {**self.filters, 'company_id': other.id})
        with self.assertRaises(ValidationError):
            self.exports._document('operations_gross', {**self.filters, 'unexpected': True})
        with self.assertRaises(ValidationError):
            self.exports._document('operations_gross', {**self.filters, 'journal_ids': ['not-an-id']})

    def test_multi_period_keeps_missing_blank_and_zero_numeric(self):
        def cell(key, label, amount, **extra):
            return {'key': key, 'label': label, 'amount': amount,
                    'negative': amount.startswith('-'), **extra}

        periods = [
            {'key': 'current', 'date_from': '2026-06-01', 'date_to': '2026-06-30',
             'display_label': 'June 2026', 'totals': [
                 cell('incoming_total', 'Sales', '1,234.50'),
                 cell('cash_total', 'Payments', '0.00'),
                 cell('difference', 'Difference', '1,234.50'),
             ], 'channels': [cell('card', 'Card', '1,234.50', percentage='100.00%')],
             'outflows': []},
            {'key': 'previous', 'date_from': '2026-05-01', 'date_to': '2026-05-31',
             'display_label': 'May 2026', 'totals': [
                 cell('incoming_total', 'Sales', '0.00'),
                 cell('cash_total', 'Payments', '0.00'),
                 cell('difference', 'Difference', '0.00'),
             ], 'channels': [], 'outflows': []},
        ]
        snapshot = {'complete': True, 'company_id': self.company.id, 'periods': periods,
                    'currency_code': self.company.currency_id.name,
                    'title': 'Sales and Payments Summary', 'basis': 'Recorded source'}
        with patch.object(type(self.report), 'get_owner_snapshot', return_value=snapshot):
            document = self.exports._document('operations_gross', self.filters)
        rows = {row['cells']['label']: row['cells'] for row in document['rows']}
        self.assertEqual(rows['Sales']['current'], Decimal('1234.50'))
        self.assertEqual(rows['Sales']['previous'], Decimal('0.00'))
        self.assertEqual(rows['Card']['current_percentage'], Decimal('1.00'))
        self.assertNotIn('previous', rows['Card'])
        self.assertNotIn('previous_percentage', rows['Card'])

    def test_normal_accounting_reader_is_allowed_and_plain_user_is_not(self):
        self._statement('40.00')
        accounting = self.env['res.users'].create({
            'name': 'Unified export accounting reader', 'login': 'unified.export.reader',
            'company_id': self.company.id, 'company_ids': [Command.set(self.company.ids)],
            'group_ids': [Command.set([
                self.env.ref('base.group_user').id,
                self.env.ref('account.group_account_user').id,
                self.env.ref('point_of_sale.group_pos_user').id,
            ])],
        })
        document = self.exports.with_user(accounting).with_context(
            allowed_company_ids=self.company.ids,
        )._document('operations_gross', self.filters)
        self.assertTrue(document['rows'])
        plain = self.env['res.users'].create({
            'name': 'Unified export plain reader', 'login': 'unified.export.plain',
            'company_id': self.company.id, 'company_ids': [Command.set(self.company.ids)],
            'group_ids': [Command.set(self.env.ref('base.group_user').ids)],
        })
        with self.assertRaises(AccessError):
            self.exports.with_user(plain).with_context(
                allowed_company_ids=self.company.ids,
            )._document('operations_gross', self.filters)


@tagged('post_install', '-at_install', 'reports_unified_export')
class TestOwnerUnifiedExportHttp(HttpCase):
    """The download route remains private, CSRF-protected and non-cacheable."""

    _statement = test_owner_cash.TestOwnerCash._statement

    def setUp(self):
        super().setUp()
        self.company = self.env.company
        self.expense = self.env['account.account'].with_company(self.company).create({
            'code': '959802', 'name': 'Unified export HTTP expense',
            'account_type': 'expense', 'company_ids': [Command.set(self.company.ids)],
        })
        self.bank = self.env['account.journal'].with_company(self.company).search([
            ('company_id', '=', self.company.id), ('type', '=', 'bank'),
        ], limit=1)
        self.assertTrue(self.bank and self.bank.default_account_id)
        self.partner = self.env['res.partner'].with_company(self.company).create({
            'name': 'Unified export HTTP partner',
        })
        self.filters = {'company_id': self.company.id, 'date_from': '2026-06-01',
                        'date_to': '2026-06-30', 'journal_ids': []}
        self._statement('40.00')

    def _csrf(self):
        # The bootstrap page carries the session-bound token used by the browser
        # download helper; this route itself is deliberately not JSON-RPC.
        page = self.url_open('/odoo')
        token = re.search(rb'''csrf_token["']?\s*:\s*["']([^"'\\]+)''', page.content)
        self.assertIsNotNone(token)
        return token.group(1).decode()

    def _download(self, csrf, file_format, filters=None):
        return self.url_open('/baseer/reports/export', data={
            'csrf_token': csrf, 'report_key': 'operations_gross',
            'file_format': file_format, 'filters': json.dumps(filters or self.filters),
        })

    def test_authenticated_downloads_are_csrf_protected_private_and_tamper_safe(self):
        self.authenticate('admin', 'admin')
        denied = self.url_open('/baseer/reports/export', data={
            'report_key': 'operations_gross', 'file_format': 'xlsx',
            'filters': json.dumps(self.filters),
        })
        self.assertNotEqual(denied.status_code, 200)
        csrf = self._csrf()
        xlsx = self._download(csrf, 'xlsx')
        self.assertEqual(xlsx.status_code, 200)
        self.assertIn('spreadsheetml.sheet', xlsx.headers['Content-Type'])
        self.assertIn('private', xlsx.headers['Cache-Control'])
        self.assertIn('no-store', xlsx.headers['Cache-Control'])
        self.assertEqual(xlsx.headers['X-Content-Type-Options'], 'nosniff')
        self.assertTrue(xlsx.content.startswith(b'PK'))
        pdf = self._download(csrf, 'pdf')
        self.assertEqual(pdf.status_code, 200)
        self.assertEqual(pdf.headers['Content-Type'], 'application/pdf')
        self.assertTrue(pdf.content.startswith(b'%PDF-'))

        foreign = self.env['res.company'].create({'name': 'Unified export HTTP foreign'})
        self.assertEqual(self._download(csrf, 'xlsx', {**self.filters,
                                                       'company_id': foreign.id}).status_code, 400)
        self.assertEqual(self._download(csrf, 'xlsx', {**self.filters,
                                                       'untrusted': True}).status_code, 400)
        self.opener.cookies.clear()
        self.assertNotEqual(self.url_open('/baseer/reports/export', data={
            'csrf_token': csrf, 'report_key': 'operations_gross', 'file_format': 'xlsx',
            'filters': json.dumps(self.filters),
        }, allow_redirects=False).status_code, 200)
