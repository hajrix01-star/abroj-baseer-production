"""Bounded, private Saudi VAT workbook; no accounting calculation lives here."""

from datetime import datetime, timezone
from decimal import Decimal
from io import BytesIO
import logging
import re
from time import monotonic

from odoo import _, api, fields, models
from odoo.exceptions import AccessError, MissingError, UserError


_logger = logging.getLogger(__name__)
_XLSX_LIMIT = 1024 * 1024
_EXPORT_SECONDS = 3
_EXCEL_NUMBER_FORMAT = '#,##0'
_ILLEGAL_XML = re.compile(r'[\x00-\x08\x0b-\x0c\x0e-\x1f]')


class BaseerTaxReportXlsx(models.TransientModel):
    _name = 'baseer.tax.report.xlsx'
    _description = 'Private Saudi VAT XLSX download'
    _transient_max_hours = 1

    company_id = fields.Many2one('res.company', required=True, readonly=True)
    period_type = fields.Selection([('month', 'Monthly'), ('quarter', 'Quarterly')], required=True, readonly=True)
    year = fields.Integer(required=True, readonly=True)
    month = fields.Char(required=True, readonly=True)
    quarter = fields.Char(required=True, readonly=True)
    display_mode = fields.Selection([('simple', 'Summary'), ('detailed', 'Detailed')], required=True, readonly=True)
    journal_ids = fields.Many2many('account.journal', string='Selected journals', readonly=True)
    # Deliberately not Binary: /web/content must never offer an alternate file route.
    file_data = fields.Text(readonly=True)
    file_name = fields.Char(readonly=True)

    @api.model
    def _download_record(self, export_id, company_id):
        """Revalidate access at GET time; the record rule also limits ORM reads."""
        if type(export_id) is not int or type(company_id) is not int:
            raise AccessError(_('Invalid VAT export reference.'))
        if company_id not in self.env.companies.ids:
            raise AccessError(_('The export company is not active.'))
        export = self.browse(export_id).exists()
        if not export:
            raise MissingError(_('The VAT export has expired. Generate it again.'))
        export.check_access('read')
        if export.create_uid.id != self.env.uid or export.company_id.id != company_id:
            raise AccessError(_('The VAT export is not available.'))
        self.env['baseer.tax.report.wizard']._hub_wizard({
            'company_id': export.company_id.id,
            'period_type': export.period_type,
            'year': export.year,
            'month': export.month,
            'quarter': export.quarter,
            'display_mode': export.display_mode,
            'journal_ids': export.journal_ids.ids,
        })
        if not export.file_data or not export.file_name:
            raise MissingError(_('The VAT export has expired. Generate it again.'))
        return export

    @api.model
    def _render_workbook(self, report, started_at=None):
        """Serialize the 16 authoritative Decimal boxes into a compact A4 workbook."""
        started_at = started_at or monotonic()
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
            from openpyxl.worksheet.page import PageMargins
        except ImportError as error:
            _logger.warning('VAT XLSX rejected: openpyxl unavailable')
            raise UserError(_('Excel export is unavailable on this server.')) from error

        company = report['company']
        currency = report['currency']
        decimals = currency.decimal_places
        if type(decimals) is not int or not 0 <= decimals <= 13:
            raise UserError(_('The company currency precision cannot be exported safely to Excel.'))
        limit = Decimal(10) ** (14 - decimals)
        number_format = _EXCEL_NUMBER_FORMAT + ('.' + '0' * decimals if decimals else '')
        rows = report['rows']
        if len(rows) != 16 or [row['number'] for row in rows] != [str(n) for n in range(1, 17)]:
            raise UserError(_('The VAT export requires the complete 16-box return.'))

        is_rtl = report['is_rtl']
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'ضريبة القيمة' if is_rtl else 'VAT Return'
        sheet.sheet_view.rightToLeft = is_rtl
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = 'C10'
        sheet.column_dimensions['A'].width = 18
        sheet.column_dimensions['B'].width = 61
        sheet.column_dimensions['C'].width = 22
        sheet.column_dimensions['D'].width = 22

        ink = '243248'
        muted = 'AAB7C4'
        negative = 'D3262E'
        separator = Side(style='hair', color='DCE3E9')
        header_fill = PatternFill('solid', fgColor='EDF1F4')
        total_fill = PatternFill('solid', fgColor='DCE0E3')
        warning_fill = PatternFill('solid', fgColor='FFF4DF')

        def literal(address, value, *, bold=False, color=ink, fill=None):
            cell = sheet[address]
            cell.value = _ILLEGAL_XML.sub('', str(value if value is not None else ''))[:32767]
            cell.data_type = 's'  # A leading =, +, -, @ must never become a formula.
            cell.hyperlink = None
            cell.font = Font(name='Arial', size=10, bold=bold, color=color)
            cell.alignment = Alignment(horizontal='right' if is_rtl else 'left', vertical='center')
            if fill:
                cell.fill = fill
            return cell

        def money(address, value, *, bold=False, fill=None):
            cell = sheet[address]
            if value is not None:
                if type(value) is not Decimal or not value.is_finite() or abs(value) >= limit:
                    _logger.warning('VAT XLSX rejected: unsafe numeric precision')
                    raise UserError(_('A VAT amount exceeds safe Excel precision for this currency.'))
                cell.value = value
            cell.number_format = number_format
            cell.font = Font(name='Arial', size=10, bold=bold,
                             color=muted if value == 0 else negative if value is not None and value < 0 else ink)
            cell.alignment = Alignment(horizontal='right', vertical='center')
            if fill:
                cell.fill = fill
            return cell

        title = 'تقرير ضريبة القيمة المضافة' if is_rtl else 'Saudi VAT Report'
        sheet.merge_cells('A1:D1')
        literal('A1', title, bold=True).font = Font(name='Arial', size=15, bold=True, color=ink)
        sheet.row_dimensions[1].height = 32
        for row_number in range(2, 6):
            sheet.row_dimensions[row_number].height = 20
        labels = {
            'company': 'الشركة' if is_rtl else 'Company',
            'currency': 'العملة' if is_rtl else 'Currency',
            'from': 'من' if is_rtl else 'From',
            'to': 'إلى' if is_rtl else 'To',
            'period': 'الفترة' if is_rtl else 'Period',
            'mode': 'العرض' if is_rtl else 'Display',
            'generated': 'وقت التوليد UTC' if is_rtl else 'Generated UTC',
            'scope': 'النطاق' if is_rtl else 'Scope',
            'journals': 'الدفاتر' if is_rtl else 'Journals',
        }
        literal('A2', labels['company'], bold=True)
        literal('B2', company.name)
        literal('C2', labels['currency'], bold=True)
        literal('D2', currency.name)
        literal('A3', labels['from'], bold=True)
        sheet['B3'] = report['date_from']
        literal('C3', labels['to'], bold=True)
        sheet['D3'] = report['date_to']
        for address in ('B3', 'D3'):
            sheet[address].number_format = 'yyyy-mm-dd'
            sheet[address].font = Font(name='Arial', size=10, color=ink)
        sheet['B3'].alignment = Alignment(horizontal='right' if is_rtl else 'left',
                                          vertical='center')
        literal('A4', labels['period'], bold=True)
        literal('B4', 'ربع سنوي' if is_rtl and report['wizard'].period_type == 'quarter' else
                'شهري' if is_rtl else report['wizard'].period_type.title())
        literal('C4', labels['mode'], bold=True)
        literal('D4', 'مفصل' if is_rtl and report['display_mode'] == 'detailed' else
                'مبسط' if is_rtl else report['display_mode'].title())
        literal('A5', labels['generated'], bold=True)
        sheet['B5'] = datetime.now(timezone.utc).replace(tzinfo=None)
        sheet['B5'].number_format = 'yyyy-mm-dd hh:mm:ss'
        sheet['B5'].font = Font(name='Arial', size=10, color=ink)
        sheet['B5'].alignment = Alignment(horizontal='right' if is_rtl else 'left',
                                          vertical='center')
        literal('C5', labels['scope'], bold=True)
        literal('D5', '16 خانة كاملة' if is_rtl else 'All 16 boxes')
        literal('A6', labels['journals'], bold=True)
        sheet.merge_cells('B6:D6')
        journal_names = report['journal_names'] or ''
        long_journals = len(journal_names) > 90
        if long_journals:
            journal_count = len(report['wizard'].journal_ids)
            summary = ('%s دفتر مختار — التفاصيل في ورقة المرشحات' % journal_count if is_rtl
                       else '%s selected journals — see Filters sheet' % journal_count)
            literal('B6', summary)
        else:
            literal('B6', journal_names or ('جميع الدفاتر' if is_rtl else 'All journals'))
        sheet.row_dimensions[6].height = 26
        if report['wizard'].journal_ids:
            sheet.merge_cells('A7:D7')
            literal('A7', 'تحليل حسب الدفاتر المختارة؛ ليس إجمالي الإقرار لجميع الدفاتر' if is_rtl
                    else 'Selected journals: analytical view, not the full VAT return',
                    bold=True, color='805A19', fill=warning_fill)
            sheet.row_dimensions[7].height = 24

        for column, label in zip('ABCD', (
            ('الخانة', 'البند', 'الأساس', 'الضريبة') if is_rtl else
            ('Box', 'Item', 'Base', 'VAT')
        )):
            literal(f'{column}9', label, bold=True, fill=header_fill)
        sheet.row_dimensions[8].height = 8
        sheet.row_dimensions[9].height = 25
        for index, row in enumerate(rows, 10):
            is_total = row['number'] in ('6', '12', '13', '16')
            fill = total_fill if is_total else None
            literal(f'A{index}', row['number'], bold=is_total, fill=fill)
            label_cell = literal(f'B{index}', row['name'].split('. ', 1)[-1],
                                 bold=is_total, fill=fill)
            label_cell.alignment = Alignment(horizontal='right' if is_rtl else 'left',
                                             vertical='center', wrap_text=True)
            money(f'C{index}', row['base'], bold=is_total, fill=fill)
            money(f'D{index}', row['tax'], bold=is_total, fill=fill)
            for cell in sheet[index]:
                cell.border = Border(bottom=separator)
            sheet.row_dimensions[index].height = (
                34 if row['number'] == '9' or len(str(row['name'])) > 45 else 23
            )

        sheet.merge_cells('A27:D27')
        literal('A27', 'قيود ضريبية غير مصنفة / Unclassified tax entries', bold=True, fill=header_fill)
        sheet.row_dimensions[27].height = 23
        exception = report['exception']
        for row_number, text_value, count, amount in (
            (28, 'قيود ضريبة قيمة مضافة غير مصنفة / Untagged VAT entries',
             exception['count'], exception['amount']),
            (29, 'قيود ضريبية أخرى غير مصنفة / Other untagged tax entries',
             exception['other_count'], exception['other_amount']),
        ):
            sheet.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=2)
            literal(f'A{row_number}', text_value)
            if type(count) is not int or count < 0:
                raise UserError(_('Invalid VAT exception count.'))
            sheet[f'C{row_number}'] = count
            sheet[f'C{row_number}'].number_format = '#,##0'
            money(f'D{row_number}', amount)
            sheet.row_dimensions[row_number].height = 23
        sheet.merge_cells('A30:D30')
        literal('A30', 'الاستثناءات خارج خانات الإقرار / Exceptions are excluded from return boxes',
                color='805A19', fill=warning_fill)
        sheet.row_dimensions[26].height = 10
        sheet.row_dimensions[30].height = 23

        sheet.print_title_rows = '1:9'
        sheet.print_area = 'A1:D30'
        sheet.print_options.horizontalCentered = True
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = 'portrait'
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.page_margins = PageMargins(left=0.35, right=0.35, top=0.5, bottom=0.5,
                                         header=0.2, footer=0.2)
        sheet.oddFooter.center.text = 'Page &P / &N'
        if long_journals:
            filters = workbook.create_sheet('المرشحات' if is_rtl else 'Filters')
            filters.sheet_view.rightToLeft = is_rtl
            filters.sheet_view.showGridLines = False
            filters.column_dimensions['A'].width = 12
            filters.column_dimensions['B'].width = 88
            filters.merge_cells('A1:B1')
            for address, value in (
                ('A1', 'الدفاتر المختارة / Selected journals'),
                ('A2', labels['company']), ('B2', company.name),
                ('A3', labels['from']), ('A4', labels['to']),
                ('A5', 'الجزء' if is_rtl else 'Part'),
                ('B5', 'أسماء الدفاتر الكاملة' if is_rtl else 'Complete journal names'),
            ):
                cell = filters[address]
                cell.value = _ILLEGAL_XML.sub('', str(value))[:32767]
                cell.data_type = 's'
                cell.hyperlink = None
                cell.font = Font(name='Arial', size=11 if address == 'A1' else 10,
                                 bold=address in ('A1', 'A5', 'B5'), color=ink)
                cell.alignment = Alignment(horizontal='right' if is_rtl else 'left', vertical='center')
            filters['B3'] = report['date_from']
            filters['B4'] = report['date_to']
            filters['B3'].number_format = filters['B4'].number_format = 'yyyy-mm-dd'
            # Split the exact source text into readable rows without losing separators.
            # Joining column B from row 6 reconstructs the full filter string.
            chunks = []
            remaining = journal_names
            while remaining:
                if len(remaining) <= 84:
                    chunks.append(remaining)
                    break
                boundary = remaining.rfind(', ', 0, 84)
                cut = boundary + 2 if boundary > 40 else 84
                chunks.append(remaining[:cut])
                remaining = remaining[cut:]
            for index, chunk in enumerate(chunks, 6):
                filters[f'A{index}'] = index - 5
                cell = filters[f'B{index}']
                cell.value = _ILLEGAL_XML.sub('', chunk)
                cell.data_type = 's'
                cell.hyperlink = None
                cell.font = Font(name='Arial', size=10, color=ink)
                cell.alignment = Alignment(horizontal='right' if is_rtl else 'left',
                                           vertical='center', wrap_text=True)
                filters.row_dimensions[index].height = 29
            filters.freeze_panes = 'B6'
            filters.print_title_rows = '1:5'
            filters.print_area = f'A1:B{5 + len(chunks)}'
            filters.sheet_properties.pageSetUpPr.fitToPage = True
            filters.page_setup.orientation = 'portrait'
            filters.page_setup.paperSize = filters.PAPERSIZE_A4
            filters.page_setup.fitToWidth = 1
            filters.page_setup.fitToHeight = 0
            filters.page_margins = PageMargins(left=0.35, right=0.35, top=0.5,
                                               bottom=0.5, header=0.2, footer=0.2)
            filters.oddFooter.center.text = 'Page &P / &N'
        output = BytesIO()
        workbook.save(output)
        payload = output.getvalue()
        duration = monotonic() - started_at
        if len(payload) > _XLSX_LIMIT or duration > _EXPORT_SECONDS:
            _logger.warning('VAT XLSX rejected: bytes=%s elapsed_ms=%s', len(payload), int(duration * 1000))
            raise UserError(_('VAT Excel export exceeded its safe size or time limit. Please retry.'))
        _logger.info('VAT XLSX prepared: bytes=%s elapsed_ms=%s', len(payload), int(duration * 1000))
        return payload
