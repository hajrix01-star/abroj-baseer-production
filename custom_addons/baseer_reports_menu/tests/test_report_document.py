"""Standalone pure contract tests: python path/to/test_report_document.py."""

from copy import deepcopy
from datetime import date, datetime
from decimal import Decimal, localcontext
import importlib.util
import base64
import re
from io import BytesIO
from pathlib import Path
import unittest
from unittest.mock import patch
from zipfile import ZipFile
from xml.etree import ElementTree

from openpyxl import load_workbook

_SPEC = importlib.util.spec_from_file_location("baseer_report_document", Path(__file__).parents[1] / "report_document.py")
engine = importlib.util.module_from_spec(_SPEC)
_SPEC.loader.exec_module(engine)


def document():
    return {
        "title": "ملخص تجريبي — Report", "company": "شركة الاختبار", "currency": "SAR", "lang": "ar", "decimals": 2,
        "metadata": [("من", date(2026, 1, 1)), ("الفلاتر", "مُرحّل فقط")],
        "columns": [{"key": "label", "label": "البند", "type": "text"},
                    {"key": "amount", "label": "المبلغ", "type": "money"},
                    {"key": "ratio", "label": "النسبة", "type": "percent"},
                    {"key": "day", "label": "التاريخ", "type": "date"}],
        "rows": [{"kind": "section", "cells": {"label": "المبيعات", "amount": Decimal("12.30")}},
                 {"kind": "detail", "indent": 1, "cells": {"label": "نقدي", "amount": Decimal("0"), "ratio": Decimal("0.125"), "day": date(2026, 1, 2)}},
                 {"kind": "detail", "cells": {"label": "خصم", "amount": Decimal("-1.20"), "ratio": None}},
                 {"kind": "total", "cells": {"label": "الإجمالي", "amount": Decimal("11.10")}}],
        "notes": ["قيم اصطناعية للاختبار فقط"],
    }


class ReportDocumentTest(unittest.TestCase):
    def test_numeric_text_ranges_keep_ltr_without_becoming_numbers(self):
        doc = document()
        doc['rows'][0]['cells']['label'] = '1–30'
        sheet = load_workbook(BytesIO(engine.render_xlsx(doc))).active
        self.assertEqual(sheet['A9'].value, '1–30')
        self.assertEqual(sheet['A9'].data_type, 's')
        self.assertEqual(sheet['A9'].alignment.readingOrder, 1)
        self.assertEqual(sheet['D10'].alignment.readingOrder, 1)
        self.assertRegex(engine.render_html(doc), r'<td class="date"[^>]*>1–30</td>')

    def test_typed_values_layout_and_no_links(self):
        doc = document()
        before = deepcopy(doc)
        data = engine.render_xlsx(doc)
        book = load_workbook(BytesIO(data), data_only=False)
        sheet = book.active
        self.assertTrue(sheet.sheet_view.rightToLeft)
        self.assertEqual(sheet.freeze_panes, "A9")
        self.assertEqual(sheet.print_title_rows, "$8:$8")
        self.assertEqual(sheet["B9"].value, 12.3)
        self.assertEqual(sheet["B10"].value, 0)
        self.assertIsNone(sheet["C9"].value)
        self.assertEqual(sheet["C10"].value, 0.125)
        self.assertEqual(sheet["B11"].value, -1.2)
        self.assertEqual(sheet["B12"].value, 11.1)
        self.assertEqual(sheet["B9"].number_format, "#,##0.00")
        self.assertEqual(sheet["B9"].font.name, "IBM Plex Sans Arabic")
        self.assertEqual(sheet["A1"].font.name, "IBM Plex Sans Arabic")
        self.assertEqual(sheet["A5"].value.date(), date(2026, 1, 1))
        self.assertEqual(sheet["A6"].value, "الفلاتر: مُرحّل فقط")
        self.assertEqual(sheet["D10"].value.date(), date(2026, 1, 2))
        for row in sheet:
            for cell in row:
                self.assertIsNone(cell.hyperlink)
                self.assertNotEqual(cell.data_type, "f")
        self.assertEqual(doc, before)
        with ZipFile(BytesIO(data)) as archive:
            xml = archive.read("xl/worksheets/sheet1.xml")
            self.assertNotIn(b"hyperlink", xml)
            root = ElementTree.fromstring(xml)
            self.assertEqual(root.findall(".//{http://schemas.openxmlformats.org/spreadsheetml/2006/main}f"), [])

    def test_literal_text_and_html_escaping(self):
        for value in ('=HYPERLINK("https://example.invalid","x")', "+SUM(A1)", "-1+2", "@SUM(A1)", '<script>alert(1)</script><img src="file:///etc/passwd">', 'https://example.invalid/?x=1&y=2'):
            doc = document()
            doc["rows"][0]["cells"]["label"] = value
            doc["title"] = value
            doc["notes"] = [value]
            sheet = load_workbook(BytesIO(engine.render_xlsx(doc))).active
            self.assertEqual(sheet["A9"].value, value)
            self.assertEqual(sheet["A9"].data_type, "s")
            html = engine.render_html(doc)
            self.assertNotIn("<script", html)
            self.assertNotIn("<img", html)
            self.assertNotIn("href=", html)
            self.assertNotIn("<a ", html)

    def test_html_layout_numbers_and_missing(self):
        html = engine.render_html(document())
        self.assertIn('lang="ar" dir="rtl"', html)
        self.assertIn("table-header-group", html)
        self.assertIn('class="detail before-total"', html)
        self.assertIn(">12.5%</td>", html)
        self.assertIn(">0.00</td>", html)
        self.assertIn(">-1.20</td>", html)
        self.assertNotIn("None", html)
        self.assertNotIn("١", html)

    def test_html_embeds_only_the_two_fixed_package_fonts(self):
        html = engine.render_html(document())
        css = html.split("<style>", 1)[1].split("</style>", 1)[0]
        fonts = re.findall(r"url\(data:font/ttf;base64,([A-Za-z0-9+/=]+)\)", css)
        self.assertEqual(len(fonts), 2)
        folder = Path(engine.__file__).parent / "static" / "fonts"
        for encoded, filename in zip(fonts, ("IBMPlexSansArabic-Regular.ttf", "IBMPlexSansArabic-Bold.ttf")):
            self.assertEqual(base64.b64decode(encoded), (folder / filename).read_bytes())
        self.assertIn("font-weight:400", css)
        self.assertIn("font-weight:700", css)
        self.assertNotIn("http:", css)
        self.assertNotIn("https:", css)
        self.assertNotIn("file:", css)
        self.assertEqual(css.count("url("), 2)

    def test_html_numeric_metadata_preserves_ltr_dates_and_escapes_other_values(self):
        doc = document()
        doc["metadata"] = [("الفترة", "2026-06-01 — 2026-06-30"),
                           ("وقت التصدير", "2026-10-23 23:33:06"),
                           ("من", date(2026, 6, 1)),
                           ("الفلاتر", '<img src="file:///private"> & all')]
        html = engine.render_html(doc)
        for value in ("2026-06-01 — 2026-06-30", "2026-10-23 23:33:06", "2026-06-01"):
            self.assertIn(f'<span class="metadata-value-ltr" dir="ltr">{value}</span>', html)
        self.assertIn('<span dir="auto">&lt;img src=&quot;file:///private&quot;&gt; &amp; all</span>', html)
        self.assertIn(".metadata-value-ltr{direction:ltr;unicode-bidi:embed;display:inline-block;}", html)
        self.assertNotIn("<img", html)

    def test_empty_and_english(self):
        doc = document()
        doc.update(lang="en", rows=[], metadata=[], notes=[])
        sheet = load_workbook(BytesIO(engine.render_xlsx(doc))).active
        self.assertFalse(sheet.sheet_view.rightToLeft)
        self.assertEqual(sheet.title, "Report")
        self.assertIn('dir="ltr"', engine.render_html(doc))

    def test_string_metadata_is_compact_literal_and_complete(self):
        doc = document()
        doc["metadata"] = [("Journals", "All journals"), ("=Unsafe label", "=1+2"),
                           ("Exported UTC", "2026-01-01 12:00:00")]
        sheet = load_workbook(BytesIO(engine.render_xlsx(doc))).active
        self.assertEqual([sheet.cell(row, 1).value for row in range(4, 7)],
                         ["Journals: All journals", "=Unsafe label: =1+2", "Exported UTC: 2026-01-01 12:00:00"])
        self.assertEqual(sheet["A5"].data_type, "s")
        self.assertIsNone(sheet["A5"].hyperlink)
        self.assertEqual(sheet.print_title_rows, "$8:$8")

    def test_invalid_number_types_and_precision(self):
        for number in (1.0, 1, True, "12", Decimal("NaN"), Decimal("Infinity"), Decimal("-Infinity"), Decimal("0.001"), Decimal("123456789012345.00"), Decimal("1E+100"), Decimal("1E-100"), Decimal("0E-100000")):
            doc = document()
            doc["rows"][0]["cells"]["amount"] = number
            with self.subTest(number=number), self.assertRaises(engine.DocumentValidationError):
                engine.validate_document(doc)
        doc = document()
        doc["rows"][0]["cells"]["amount"] = Decimal("9999999999999.99")
        engine.validate_document(doc)
        with localcontext() as context:
            context.prec = 3
            engine.validate_document(doc)
            self.assertIn("9,999,999,999,999.99", engine.render_html(doc))

    def test_percentage_precision(self):
        doc = document()
        for value in (Decimal("0.1234567890123456"), Decimal("1E-16")):
            doc["rows"][1]["cells"]["ratio"] = value
            with self.assertRaises(engine.DocumentValidationError):
                engine.validate_document(doc)
        for value, formatted in (("0", "0%"), ("1.00", "100%"), ("0.000001", "0.0001%"), ("-0.125", "-12.5%")):
            doc["rows"][1]["cells"]["ratio"] = Decimal(value)
            self.assertIn(">" + formatted + "</td>", engine.render_html(doc))

    def test_malformed_and_oversized_documents(self):
        mutations = [lambda d: d.update(lang="fr"), lambda d: d.update(decimals=True), lambda d: d.update(decimals=3),
                     lambda d: d.update(title="x" * 257), lambda d: d.update(company=""),
                     lambda d: d.update(notes=["x\x00"]), lambda d: d.update(href="bad"),
                     lambda d: d.update(columns=[]), lambda d: d["columns"].append(d["columns"][0]),
                     lambda d: d["columns"][0].update(width=5), lambda d: d["columns"][0].update(width=20.0),
                     lambda d: d["rows"][0].update(indent=4), lambda d: d["rows"][0].update(kind="evil"),
                     lambda d: d["rows"][0]["cells"].update(unknown="x"),
                     lambda d: d["rows"][0]["cells"].update(label="x" * 4097),
                     lambda d: d["rows"][0]["cells"].update(day=datetime(2026, 1, 1)),
                     lambda d: d.update(rows=[d["rows"][0]] * 5001)]
        for mutate in mutations:
            doc = document()
            mutate(doc)
            with self.subTest(mutate=mutate), self.assertRaises(engine.DocumentValidationError):
                engine.validate_document(doc)
        doc = document()
        doc["columns"] = [{"key": "c" + str(i), "label": "Test", "type": "text"} for i in range(11)]
        doc["rows"] = [{"kind": "detail", "cells": {}}] * 5000
        with self.assertRaises(engine.DocumentValidationError):
            engine.validate_document(doc)

    def test_output_caps(self):
        with patch.object(engine, "MAX_OUTPUT_BYTES", 1024):
            with self.assertRaises(engine.DocumentValidationError):
                engine.render_html(document())
            with self.assertRaises(engine.DocumentValidationError):
                engine.render_xlsx(document())
            doc = document()
            doc["rows"] = [{"kind": "detail", "cells": {"label": "x" * 300}} for _ in range(4)]
            with self.assertRaises(engine.DocumentValidationError):
                engine.validate_document(doc)

    def test_long_text_wraps_or_fails_without_clipping(self):
        doc = document()
        doc["rows"][0]["cells"]["label"] = "وصف طويل " * 10
        sheet = load_workbook(BytesIO(engine.render_xlsx(doc))).active
        self.assertGreater(sheet.row_dimensions[9].height, 20)
        doc["rows"][0]["cells"]["label"] = "ط" * 4096
        with self.assertRaises(engine.DocumentValidationError):
            engine.render_xlsx(doc)

    def test_bounded_representative_fixture(self):
        doc = document()
        doc["rows"] = [deepcopy(doc["rows"][1]) for _ in range(5000)]
        self.assertLess(len(engine.render_xlsx(doc)), engine.MAX_OUTPUT_BYTES)
        self.assertLess(len(engine.render_html(doc).encode("utf-8")), engine.MAX_OUTPUT_BYTES)


if __name__ == "__main__":
    unittest.main()
